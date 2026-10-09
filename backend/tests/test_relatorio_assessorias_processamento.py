"""Fase 8 (Gerador de Relatórios Mensais das Assessorias) — processamento.py.

Ponta a ponta com planilhas pequenas construídas em memória (fictícias,
salvas em arquivo temporário só porque `processar_upload` recebe caminho
de arquivo, não Workbook — mesma interface que o upload real usa)."""
from __future__ import annotations

import tempfile
from datetime import date

from openpyxl import Workbook

from app.auth import hash_senha
from app.models import EmpresaCliente, Usuario
from app.relatorio_assessorias import armazenamento, processamento


def _salvar(wb: Workbook) -> str:
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        wb.save(tmp.name)
        return tmp.name


def _workbook_laudo():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("AGO26")
    ws.append(["DATA", "EMPRESA", "TIPO DE LAUDO", "ENTRADA DE LAUDO", "LAUDO PRONTO", "DATA", "ENTREGUE A EMPRESA"])
    ws.append(["01/08/2026", "TESTE SA", "AUTO", "", "Sim", "05/08/2026", "Sim"])
    return _salvar(wb)


def _workbook_iniciais():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("AGOSTO 2026")
    ws.append(["DATA DE RECEBIMENTO", "ASSESSORIA", "DATA DE DISTRIBUIÇÃO", "PROTOCOLADO", "ESTADOS", "NÚMERO DE PROCESSO"])
    ws.append(["03/08/2026", "TESTE SA", "06/08/2026", "SIM", "RJ", "0001"])
    return _salvar(wb)


def _workbook_agendamento():
    wb = Workbook()
    wb.remove(wb.active)
    ws_judicial = wb.create_sheet("JUDICIAL")
    ws_judicial.append([None, "NOME DO CLIENTE", "PROCESSO", "EMPRESA"])
    ws_judicial.append(["01/03/2026", "Fulano", "0001", "TESTE SA"])
    ws_pauta = wb.create_sheet("PAUTA DA SEMANA CONTRARIA")
    ws_pauta.append(["PAUTA DA SEMANA CONTRARIA"])
    ws_pauta.append(["DATA", "NOME DO CLIENTE", "ESTADO", "PROCESSO", "EMPRESA"])
    ws_pauta.append(["01/04/2026", "Beltrano", "SP", "0002", "TESTE SA"])
    return _salvar(wb)


def _workbook_extrajudicial():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("AGENDAMENTO")
    ws.append(["DATA DE RECEBIMENTO", "EMPRESA", "NOME COMPLETO", "DATA DE AGENDAMENTO", "PRESENÇA"])
    ws.append(["05/08/2026", "TESTE SA", "Ciclano", "20/08/2026", "PRESENTE"])
    return _salvar(wb)


def _workbook_contrarias():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("TESTE SA")
    ws.append(["TESTE SA"])
    ws.append(["título"])
    ws.append(["DATA DE RECEBIMENTO", "NOME DO AUTOR", "Nº DO PROCESSO", "ESTADO", "VALOR DA CAUSA", "STATUS ATUAL", "TIPO AÇÃO"])
    ws.append(["10/08/2026", "Fulano de Tal", "0813055-21.2026.8.10.0001", "RJ", "R$ 1.000,00", "ATIVO", "JUDICIAL"])
    return _salvar(wb)


def _workbook_procon():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("TESTE SA")
    ws.append(["TESTE SA"])
    ws.append(["DATA DE RECEBIMENTO", "NOME DO AUTOR", "Nº DO PROCESSO", "ESTADO", "STATUS ATUAL", "TIPO AÇÃO"])
    ws.append(["10/08/2026", "Beltrano", "0002", "SP", "ATIVO", "PROCON"])
    return _salvar(wb)


def test_processar_upload_salva_um_resultado_por_assessoria(db):
    usuario = Usuario(nome="Teste", email="proc@teste.local", senha_hash=hash_senha("x"), papel_global="ADMIN_SUPERIOR")
    db.add(usuario)
    db.add(EmpresaCliente(nome="TESTE SA"))
    db.commit()

    caminhos = {
        "laudo": _workbook_laudo(),
        "iniciais": _workbook_iniciais(),
        "agendamento": _workbook_agendamento(),
        "extrajudicial": _workbook_extrajudicial(),
        "contrarias": _workbook_contrarias(),
        "procon": _workbook_procon(),
    }

    resultados, desconhecidas = processamento.processar_upload(
        db, caminhos, mes=8, ano=2026, data_corte=date(2026, 9, 4), usuario=usuario
    )

    assert len(resultados) == 1
    assert resultados[0].assessoria == "TESTE SA"
    dados = armazenamento.dados_efetivos(resultados[0])
    assert dados.laudos.elaborados == 1
    assert dados.iniciais.distribuidos_no_mes == 1
    assert dados.extrajudiciais.enviadas == 1
    assert dados.audiencias_judiciais.quantidade == 1
    assert dados.audiencias_contrarias.quantidade == 1
    assert dados.contrarias.ativos_total == 1
    assert len(dados.procons.lista) == 1
    assert desconhecidas == []


def test_processar_upload_assessoria_sem_movimento_fica_zerada(db):
    usuario = Usuario(nome="Teste", email="proc2@teste.local", senha_hash=hash_senha("x"), papel_global="ADMIN_SUPERIOR")
    db.add(usuario)
    db.add(EmpresaCliente(nome="TESTE SA"))
    db.add(EmpresaCliente(nome="OUTRA SEM MOVIMENTO"))
    db.commit()

    caminhos = {
        "laudo": _workbook_laudo(),
        "iniciais": _workbook_iniciais(),
        "agendamento": _workbook_agendamento(),
        "extrajudicial": _workbook_extrajudicial(),
        "contrarias": _workbook_contrarias(),
        "procon": _workbook_procon(),
    }

    resultados, _ = processamento.processar_upload(
        db, caminhos, mes=8, ano=2026, data_corte=date(2026, 9, 4), usuario=usuario
    )

    assert len(resultados) == 2
    sem_movimento = next(r for r in resultados if r.assessoria == "OUTRA SEM MOVIMENTO")
    dados = armazenamento.dados_efetivos(sem_movimento)
    assert dados.laudos.elaborados == 0
    assert dados.contrarias.ativos_total == 0
