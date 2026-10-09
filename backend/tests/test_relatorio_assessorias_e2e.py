"""Fase 9 — teste de ponta a ponta contra o critério de aceite da
especificação: relatório EWS de agosto/2026, planilhas atuais, data de
corte 04/09/2026 (página 11, "Critérios de aceite e testes").

Só roda se `dados_locais_nao_versionados/` existir (planilhas reais com
dado pessoal — nunca committadas, ver DECISIONS.md 2026-10-08 e
.gitignore); em qualquer outro ambiente (CI, outra máquina) este teste é
pulado — não é um teste "quebrado", é um teste que não tem como rodar sem
os dados reais da Clara.

Um valor não é igual ao da tabela da especificação por decisão explícita,
não por engano: "Audiências judiciais (acumulado no ano)" — a
especificação lista 11, mas a investigação da Fase 4 (ver DECISIONS.md
2026-10-09, "Fase 4... Investigado e resolvido") concluiu que 10 é o
número correto pra esses dados — a Clara confirmou que a 11ª audiência do
relatório original provavelmente foi um acréscimo manual, não algo que a
planilha capture. Testado contra 10, não contra 11."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from app.auth import hash_senha
from app.models import EmpresaCliente, Usuario
from app.relatorio_assessorias import armazenamento, processamento
from app.relatorio_assessorias.render import renderizar_docx

_PASTA_DADOS_REAIS = Path(__file__).resolve().parent.parent.parent / "dados_locais_nao_versionados"

pytestmark = pytest.mark.skipif(
    not _PASTA_DADOS_REAIS.is_dir(),
    reason="planilhas reais não disponíveis neste ambiente (dados_locais_nao_versionados/ não existe)",
)


def _caminho(nome: str) -> str:
    return str(_PASTA_DADOS_REAIS / nome)


@pytest.fixture()
def caminhos_reais() -> dict[str, str]:
    return {
        "laudo": _caminho("PLANILHA LAUDO - ELITE (2).xlsx"),
        "iniciais": _caminho("PLANILHA DE INICIAS - PROTOCOLOS v2.xlsx"),
        "agendamento": _caminho("PLANILHA 2026 (1).xlsx"),
        "extrajudicial": _caminho("AGENDAMENTO (10).xlsx"),
        "contrarias": _caminho("CONTRÁRIAS - NOVO.xlsx"),
        "procon": _caminho("PROCON E EXTRAJUDICIAL 2026.xlsx"),
    }


def test_criterio_de_aceite_ews_agosto_2026(db, caminhos_reais):
    usuario = Usuario(
        nome="Teste E2E", email="e2e.relatorio@teste.local", senha_hash=hash_senha("x"), papel_global="ADMIN_SUPERIOR"
    )
    db.add(usuario)
    db.add(EmpresaCliente(nome="EWS"))
    db.commit()

    resultados, _desconhecidas = processamento.processar_upload(
        db, caminhos_reais, mes=8, ano=2026, data_corte=date(2026, 9, 4), usuario=usuario
    )

    registro_ews = next(r for r in resultados if r.assessoria == "EWS")
    dados = armazenamento.dados_efetivos(registro_ews)

    # Laudos elaborados no mês
    assert dados.laudos.elaborados == 14

    # Audiências extrajudiciais
    assert dados.extrajudiciais.enviadas == 16
    assert dados.extrajudiciais.realizadas == 10
    assert dados.extrajudiciais.pendentes_proximos_meses == 19
    assert [c.nome for c in dados.extrajudiciais.clientes_ausentes] == ["ANTONIO BESERRA DA COSTA"]

    # Audiências judiciais (acumulado no ano) — 10, não 11 (ver docstring do módulo)
    assert dados.audiencias_judiciais.quantidade == 10

    # Audiências de processos contrários (acumulado no ano)
    assert dados.audiencias_contrarias.quantidade == 5

    # Ações contrárias
    assert len(dados.contrarias.lista_judiciais) == 27
    assert len(dados.contrarias.lista_trabalhistas) == 2
    assert dados.contrarias.ativos_total == 29
    assert dados.contrarias.ativos_por_uf.get("RJ") == 3

    # Iniciais — processos distribuídos no mês
    assert dados.iniciais.distribuidos_no_mes == 2

    # Render: o .docx final abre e não sobra marcador nenhum
    conteudo = renderizar_docx(dados)
    doc = Document(BytesIO(conteudo))
    texto_total = "\n".join(p.text for p in doc.paragraphs)
    for tabela in doc.tables:
        for linha in tabela.rows:
            texto_total += "\n" + " | ".join(c.text for c in linha.cells)
    assert "{{" not in texto_total
    assert "{%" not in texto_total
