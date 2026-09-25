import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from sqlalchemy import event

from app.models import EventoProcesso, Processo
from app.services.empresas import get_or_create_empresa
from app.services.processos import (
    _mes_referencia_da_aba,
    _pessoa_valida,
    _separar_assessoria,
    apagar_todos_processos,
    gerar_relatorio_geral,
    gerar_relatorio_por_empresa,
    importar_planilha,
    marcar_resolvido,
    status_prazo,
)


def _processo(db, numero="5012298-14.2025.8.13.0231", advogada="DRA KELLY", assistente="DANILO", empresa="ABSOLUTA"):
    empresa_obj = get_or_create_empresa(db, empresa)
    processo = Processo(
        numero_processo=numero,
        empresa_cliente_id=empresa_obj.id,
        nome_cliente="Fulano de Tal",
        advogada=advogada,
        assistente=assistente,
    )
    db.add(processo)
    db.commit()
    return processo


def test_status_prazo_pendente_quando_ainda_nao_venceu(db):
    processo = _processo(db)
    evento = EventoProcesso(
        processo_id=processo.id, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS",
        prazo_fatal=True, data_prazo=date.today() + timedelta(days=5),
    )
    assert status_prazo(evento) == "PENDENTE"


def test_status_prazo_perdido_quando_vencido_sem_resolver(db):
    processo = _processo(db)
    evento = EventoProcesso(
        processo_id=processo.id, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS",
        prazo_fatal=True, data_prazo=date.today() - timedelta(days=2),
    )
    assert status_prazo(evento) == "PERDIDO"


def test_status_prazo_cumprido_quando_resolvido_antes(db):
    processo = _processo(db)
    prazo = date.today() + timedelta(days=3)
    evento = EventoProcesso(
        processo_id=processo.id, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS",
        prazo_fatal=True, data_prazo=prazo, resolvido=True,
        resolvido_em=datetime.combine(prazo - timedelta(days=1), datetime.min.time()),
    )
    assert status_prazo(evento) == "CUMPRIDO"


def test_status_prazo_cumprido_com_atraso(db):
    processo = _processo(db)
    prazo = date.today() - timedelta(days=1)
    evento = EventoProcesso(
        processo_id=processo.id, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS",
        prazo_fatal=True, data_prazo=prazo, resolvido=True,
        resolvido_em=datetime.combine(date.today(), datetime.min.time()),
    )
    assert status_prazo(evento) == "CUMPRIDO_COM_ATRASO"


def test_status_prazo_none_sem_data_prazo():
    from app.models import EventoProcesso as EP

    evento = EP(processo_id=1, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS", prazo_fatal=True)
    assert status_prazo(evento) is None


def test_marcar_resolvido(db):
    processo = _processo(db)
    evento = EventoProcesso(
        processo_id=processo.id, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS",
        prazo_fatal=True, data_prazo=date.today() + timedelta(days=1),
    )
    db.add(evento)
    db.commit()

    resolvido = marcar_resolvido(db, evento.id)
    assert resolvido.resolvido is True
    assert resolvido.resolvido_em is not None


def test_import_prazo_fatal_so_quando_coluna_e_sim(db, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO", "PRAZO FATAL"])
    numeros = [
        "1111111-11.2026.8.11.0001",
        "2222222-22.2026.8.11.0002",
        "3333333-33.2026.8.11.0003",
        "4444444-44.2026.8.11.0004",
    ]
    valores_prazo = ["SIM", "sim ", "NÃO", ""]
    for numero, valor in zip(numeros, valores_prazo, strict=True):
        ws.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS", valor])
    path = Path(tempfile.mkdtemp()) / "processos.xlsx"
    wb.save(path)

    importar_planilha(db, str(path))

    eventos = {e.processo.numero_processo: e for e in db.query(EventoProcesso).all()}
    assert eventos[numeros[0]].prazo_fatal is True  # "SIM"
    assert eventos[numeros[1]].prazo_fatal is True  # "sim " — normalizado
    assert eventos[numeros[2]].prazo_fatal is False  # "NÃO" não é fatal
    assert eventos[numeros[3]].prazo_fatal is False  # vazio não é fatal


def test_import_preenche_mes_referencia_a_partir_do_nome_da_aba(db):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SETEMBRO26"
    ws.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO"])
    numero = "5555555-55.2026.8.11.0005"
    ws.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS"])
    path = Path(tempfile.mkdtemp()) / "processos.xlsx"
    wb.save(path)

    importar_planilha(db, str(path))

    evento = db.query(EventoProcesso).join(Processo).filter(Processo.numero_processo == numero).one()
    assert evento.mes_referencia == "SETEMBRO/2026"


def test_mes_referencia_da_aba():
    assert _mes_referencia_da_aba("SETEMBRO26") == "SETEMBRO/2026"
    assert _mes_referencia_da_aba("Setembro 2026") == "SETEMBRO/2026"
    assert _mes_referencia_da_aba("OUTUBRO-26") == "OUTUBRO/2026"
    assert _mes_referencia_da_aba("MARÇO") == "MARCO"
    assert _mes_referencia_da_aba("DOCS E CUSTAS") is None
    assert _mes_referencia_da_aba("FATAL") is None


def test_mes_referencia_da_aba_aceita_abreviacao_de_3_letras():
    assert _mes_referencia_da_aba("JUN-25") == "JUNHO/2025"
    assert _mes_referencia_da_aba("JUL-25") == "JULHO/2025"
    assert _mes_referencia_da_aba("AGO-25") == "AGOSTO/2025"
    assert _mes_referencia_da_aba("SET- 25") == "SETEMBRO/2025"
    assert _mes_referencia_da_aba("OUT- 25") == "OUTUBRO/2025"
    assert _mes_referencia_da_aba("MAR") == "MARCO"


def test_pessoa_valida_rejeita_valores_parecidos_com_data():
    assert _pessoa_valida("DANILO") == "DANILO"
    assert _pessoa_valida("2025-12-03 00:00:00") is None
    assert _pessoa_valida("2025-12-03") is None
    assert _pessoa_valida("03/12/2025") is None
    assert _pessoa_valida("  ") is None


def test_separar_assessoria():
    assert _separar_assessoria("HUNTING - Fulana de Tal (CONTR. Beltrano)") == "HUNTING"
    assert _separar_assessoria("DRA KELLY") is None  # sem "-", não tem assessoria
    assert _separar_assessoria(None) is None
    assert _separar_assessoria("") is None


def test_import_extrai_assessoria_da_advogada_sem_alterar_advogada(db):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO", "ADVOGADA"])
    numero = "6666666-66.2026.8.11.0006"
    ws.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS", "HUNTING - Fulana de Tal"])
    path = Path(tempfile.mkdtemp()) / "processos.xlsx"
    wb.save(path)

    importar_planilha(db, str(path))

    processo = db.query(Processo).filter_by(numero_processo=numero).one()
    assert processo.advogada == "HUNTING - Fulana de Tal"
    assert processo.assessoria == "HUNTING"


def test_import_reimportacao_atualiza_evento_existente_com_dado_novo(db, tmp_path):
    numero = "9999999-99.2026.8.11.0009"

    def _planilha(nome_arquivo, observacao, prazo_fatal):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO", "PRAZO FATAL", "OBSERVAÇÃO"])
        ws.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS", prazo_fatal, observacao])
        caminho = tmp_path / nome_arquivo
        wb.save(caminho)
        return str(caminho)

    resumo1 = importar_planilha(db, _planilha("v1.xlsx", "observação original", ""))
    assert resumo1.linhas_novas == 1
    assert resumo1.linhas_atualizadas == 0

    evento = db.query(EventoProcesso).join(Processo).filter(Processo.numero_processo == numero).one()
    assert evento.observacao == "observação original"
    assert evento.prazo_fatal is False

    # Reimporta o mesmo andamento (mesmo processo+data+evento), agora com
    # observação corrigida e PRAZO FATAL marcado — a Clara pediu que isso
    # atualize o registro em vez de ser ignorado como "já existente".
    resumo2 = importar_planilha(db, _planilha("v2.xlsx", "observação corrigida", "SIM"))
    assert resumo2.linhas_novas == 0
    assert resumo2.linhas_ja_existentes == 1
    assert resumo2.linhas_atualizadas == 1

    db.refresh(evento)
    assert evento.observacao == "observação corrigida"
    assert evento.prazo_fatal is True


def test_import_reimportacao_nao_apaga_dado_quando_planilha_vem_vazia(db, tmp_path):
    numero = "1010101-01.2026.8.11.0010"

    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO", "OBSERVAÇÃO"])
    ws1.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS", "observação importante"])
    path1 = tmp_path / "v1.xlsx"
    wb1.save(path1)
    importar_planilha(db, str(path1))

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO", "OBSERVAÇÃO"])
    ws2.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS", ""])
    path2 = tmp_path / "v2.xlsx"
    wb2.save(path2)
    resumo2 = importar_planilha(db, str(path2))
    assert resumo2.linhas_atualizadas == 0  # nada mudou — observação vazia não sobrescreve

    evento = db.query(EventoProcesso).join(Processo).filter(Processo.numero_processo == numero).one()
    assert evento.observacao == "observação importante"


def test_import_reimportacao_nao_mexe_em_resolvido(db, tmp_path):
    numero = "1212121-21.2026.8.11.0011"
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO"])
    ws1.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS"])
    path1 = tmp_path / "v1.xlsx"
    wb1.save(path1)
    importar_planilha(db, str(path1))

    evento = db.query(EventoProcesso).join(Processo).filter(Processo.numero_processo == numero).one()
    marcar_resolvido(db, evento.id)
    assert evento.resolvido is True

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO", "OBSERVAÇÃO"])
    ws2.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS", "nova observação"])
    path2 = tmp_path / "v2.xlsx"
    wb2.save(path2)
    importar_planilha(db, str(path2))

    db.refresh(evento)
    assert evento.resolvido is True  # reimport nunca desfaz isso
    assert evento.observacao == "nova observação"  # mas outros campos seguem atualizando


def test_import_reimportacao_atualiza_nome_cliente_do_processo(db, tmp_path):
    numero = "1313131-31.2026.8.11.0012"
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO"])
    ws1.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS"])
    path1 = tmp_path / "v1.xlsx"
    wb1.save(path1)
    importar_planilha(db, str(path1))

    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO"])
    ws2.append(["ABSOLUTA - Fulano de Tal Corrigido", numero, date(2026, 9, 2), "DOCUMENTOS"])
    path2 = tmp_path / "v2.xlsx"
    wb2.save(path2)
    importar_planilha(db, str(path2))

    processo = db.query(Processo).filter_by(numero_processo=numero).one()
    assert processo.nome_cliente == "Fulano de Tal Corrigido"


def test_import_varios_eventos_do_mesmo_processo_novo_na_mesma_planilha(db, tmp_path):
    """Regressão de performance: o import parou de dar `db.flush()` logo
    depois de criar um `Processo` novo (fazia isso pra cada processo novo —
    ~1/3 do tempo total de um import de ~45 mil linhas, medido localmente).
    Isso só funciona se múltiplos eventos de um MESMO processo, ainda sem
    `.id` (não foi pro banco ainda), conseguirem se ligar a ele corretamente
    via o relacionamento do SQLAlchemy — em vez de `processo_id=processo.id`
    (que seria `None` nesse momento). Esse teste cobre exatamente esse caso:
    dois eventos do mesmo processo novo, na mesma planilha, no mesmo import."""
    numero = "1414141-41.2026.8.11.0013"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["CLIENTE", "Nº PROCESSO", "DATA", "EVENTO"])
    ws.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 1), "CUSTAS"])
    ws.append(["ABSOLUTA - Fulano de Tal", numero, date(2026, 9, 10), "DOCUMENTOS"])
    path = tmp_path / "processos.xlsx"
    wb.save(path)

    resumo = importar_planilha(db, str(path))
    assert resumo.linhas_novas == 2

    processos = db.query(Processo).filter_by(numero_processo=numero).all()
    assert len(processos) == 1  # não duplicou o processo
    eventos = db.query(EventoProcesso).filter_by(processo_id=processos[0].id).all()
    assert {e.tipo_evento_nome for e in eventos} == {"CUSTAS", "DOCUMENTOS"}


def test_apagar_todos_processos_remove_processos_e_eventos_e_devolve_a_contagem(db):
    p1 = _processo(db, numero="1717171-71.2026.8.11.0016")
    p2 = _processo(db, numero="1818181-81.2026.8.11.0017")
    db.add_all(
        [
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS"),
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 2), tipo_evento_nome="DOCUMENTOS"),
            EventoProcesso(processo_id=p2.id, data=date(2026, 9, 1), tipo_evento_nome="CUSTAS"),
        ]
    )
    db.commit()

    apagados = apagar_todos_processos(db)

    assert apagados == 2
    assert db.query(Processo).count() == 0
    assert db.query(EventoProcesso).count() == 0
    # não mexe nas empresas-clientes
    assert get_or_create_empresa(db, "ABSOLUTA").nome == "ABSOLUTA"


def test_apagar_todos_processos_com_banco_ja_vazio(db):
    assert apagar_todos_processos(db) == 0


def test_import_reconhece_empresa_em_coluna_propria_alem_do_formato_com_hifen(db):
    """Bug real reportado pela Clara (2026-09-24): abas mais recentes (ex.:
    SETEMBRO26) passaram a ter EMPRESA numa coluna própria, sem o formato
    antigo "EMPRESA - Cliente" embutido em CLIENTE. Sem esse reconhecimento,
    quase toda a aba era descartada do import por "empresa não reconhecida"
    — o relatório do mês corrente ficava com quase nada."""
    numero = "1919191-91.2026.8.11.0018"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SETEMBRO26"
    ws.append(["ASSISTENTE", "MÊS", "DIA", "EMPRESA", "CLIENTE", "Nº PROCESSO", "EVENTO"])
    ws.append(["DANILO", "SETEMBRO", date(2026, 9, 5), "ABSOLUTA", "Fulano de Tal", numero, "CUSTAS"])
    path = Path(tempfile.mkdtemp()) / "processos.xlsx"
    wb.save(path)

    resumo = importar_planilha(db, str(path))
    assert resumo.linhas_novas == 1
    assert resumo.linhas_sem_empresa_reconhecida == 0

    processo = db.query(Processo).filter_by(numero_processo=numero).one()
    assert processo.nome_cliente == "Fulano de Tal"
    assert processo.empresa_cliente_id == get_or_create_empresa(db, "ABSOLUTA").id


def test_import_marca_data_de_liberacao_e_relatorio_a_exclui_do_periodo(db):
    """Bug real (2026-09-24): abas "coringa" (fatais, Dra Galzo, DOCS E
    CUSTAS etc.) não têm data de andamento real — só a data em que a Dra
    inseriu o cliente na planilha. Essas linhas não podem contar como se
    fossem eventos daquele mês no relatório por período (a Clara
    confirmou), mas continuam existindo no sistema normalmente."""
    numero = "2020202-02.2026.8.11.0019"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dra Teste"
    ws.append(["ASSISTENTE", "ANO", "MÊS", "DATA DE LIBERAÇÃO - QUANDO A DRA INSERIIU O CLIENTE NA PLANILHA", "CLIENTE", "Nº PROCESSO", "EVENTO"])
    ws.append(["DANILO", 2026, "SETEMBRO", date(2026, 9, 5), "ABSOLUTA - Fulano de Tal", numero, "CUSTAS"])
    path = Path(tempfile.mkdtemp()) / "processos.xlsx"
    wb.save(path)

    resumo = importar_planilha(db, str(path))
    assert resumo.linhas_novas == 1

    evento = db.query(EventoProcesso).join(Processo).filter(Processo.numero_processo == numero).one()
    assert evento.data == date(2026, 9, 5)
    assert evento.data_e_liberacao is True

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    assert relatorio.secoes == []  # excluído do relatório por período


def test_import_data_real_de_aba_dia_nao_e_marcada_como_liberacao(db):
    """Regressão do teste acima: uma aba cuja coluna de data se chama "DIA"
    (ex.: FATAIS 08 na planilha real) tem data de andamento de verdade, não
    é a mesma coisa que "DATA DE LIBERAÇÃO" — não pode ser excluída do
    relatório por período."""
    numero = "2121212-12.2026.8.11.0020"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FATAIS 09"
    ws.append(["ASSISTENTE", "MÊS", "DIA", "CLIENTE", "Nº PROCESSO", "EVENTO"])
    ws.append(["DANILO", "SETEMBRO", date(2026, 9, 5), "ABSOLUTA - Fulano de Tal", numero, "CUSTAS"])
    path = Path(tempfile.mkdtemp()) / "processos.xlsx"
    wb.save(path)

    importar_planilha(db, str(path))

    evento = db.query(EventoProcesso).join(Processo).filter(Processo.numero_processo == numero).one()
    assert evento.data_e_liberacao is False

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    assert relatorio.secoes[0].total_processos == 1


# --- Bloco 1 (2026-09-25): relatório Geral / Por empresa ------------------


def test_relatorio_geral_uma_secao_por_empresa_ordenadas_alfabeticamente(db):
    p1 = _processo(db, numero="3030303-30.2026.8.11.0030", assistente="DANILO", empresa="ZENITH")
    p2 = _processo(db, numero="4040404-40.2026.8.11.0040", assistente="DANILO", empresa="ABSOLUTA")
    db.add_all(
        [
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
            EventoProcesso(processo_id=p2.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
        ]
    )
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    assert [s.empresa for s in relatorio.secoes] == ["ABSOLUTA", "ZENITH"]


def test_relatorio_geral_total_bate_com_numero_de_linhas(db):
    p1 = _processo(db, numero="5050505-50.2026.8.11.0050", assistente="DANILO")
    p2 = _processo(db, numero="6060606-60.2026.8.11.0060", assistente="JULIA")
    db.add_all(
        [
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 6), tipo_evento_nome="DOCUMENTOS"),  # mesmo processo, 2º andamento
            EventoProcesso(processo_id=p2.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
        ]
    )
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    secao = relatorio.secoes[0]
    assert secao.total_processos == 2  # 2 processos, não 3 andamentos
    assert len(secao.linhas) == 2  # Parte 1: uma linha por processo
    assert len(secao.linhas_resumo) == 2  # Parte 2: idem


def test_relatorio_geral_evento_e_fatal_sao_do_andamento_mais_recente(db):
    """"Evento"/"Fatal" mostrados são do andamento mais recente do
    processo, por `criado_em` (data/hora de registro) — a pedido explícito
    da Clara, não pela ordem de inserção nem pela `data` do andamento."""
    processo = _processo(db, numero="7070707-70.2026.8.11.0070", assistente="DANILO")
    db.add_all(
        [
            EventoProcesso(
                processo_id=processo.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS",
                prazo_fatal=False, criado_em=datetime.combine(date(2026, 9, 5), datetime.min.time()),
            ),
            EventoProcesso(
                processo_id=processo.id, data=date(2026, 9, 1), tipo_evento_nome="DOCUMENTOS",
                # registrado um dia depois, mesmo com `data` (do andamento) menor
                prazo_fatal=True, criado_em=datetime.combine(date(2026, 9, 6), datetime.min.time()),
            ),
        ]
    )
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    linha = relatorio.secoes[0].linhas[0]
    assert linha.evento == "DOCUMENTOS"
    assert linha.fatal is True


def test_relatorio_geral_fatal_sim_nao(db):
    processo = _processo(db, numero="8080808-80.2026.8.11.0080", assistente="DANILO")
    db.add(EventoProcesso(processo_id=processo.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS", prazo_fatal=True))
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    assert relatorio.secoes[0].linhas[0].fatal is True


def test_relatorio_geral_parte2_so_tem_cliente_processo_evento_observacao(db):
    processo = _processo(db, numero="9090909-90.2026.8.11.0090", assistente="DANILO")
    db.add(
        EventoProcesso(
            processo_id=processo.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS", observacao="obs teste",
        )
    )
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    linha = relatorio.secoes[0].linhas_resumo[0]
    assert linha.cliente == "Fulano de Tal"
    assert linha.numero_processo == "9090909-90.2026.8.11.0090"
    assert linha.evento == "CUSTAS"
    assert linha.observacao == "obs teste"
    assert not hasattr(linha, "assistente")
    assert not hasattr(linha, "fatal")


def test_relatorio_geral_sem_observacao_mostra_travessao(db):
    processo = _processo(db, numero="1234567-89.2026.8.11.0091", assistente="DANILO")
    db.add(EventoProcesso(processo_id=processo.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"))
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    assert relatorio.secoes[0].linhas_resumo[0].observacao == "—"


def test_relatorio_geral_filtra_por_assistente_em_todas_as_empresas(db):
    p1 = _processo(db, numero="2345678-90.2026.8.11.0092", assistente="DANILO", empresa="ABSOLUTA")
    p2 = _processo(db, numero="3456789-01.2026.8.11.0093", assistente="JULIA", empresa="ZENITH")
    db.add_all(
        [
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
            EventoProcesso(processo_id=p2.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
        ]
    )
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30), filtro_assistente="danilo")
    assert [s.empresa for s in relatorio.secoes] == ["ABSOLUTA"]


def test_relatorio_geral_empresa_sem_processo_no_periodo_nao_aparece(db):
    empresa_vazia = get_or_create_empresa(db, "SEM PROCESSO NO PERIODO")
    db.add(
        Processo(numero_processo="4567890-12.2026.8.11.0094", empresa_cliente_id=empresa_vazia.id, nome_cliente="X")
    )
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    assert relatorio.secoes == []


def test_relatorio_geral_ordena_por_assistente_depois_numero_processo(db):
    empresa = get_or_create_empresa(db, "ABSOLUTA")
    p_julia = Processo(numero_processo="9999999-99.2026.8.11.0099", empresa_cliente_id=empresa.id, nome_cliente="Z", assistente="JULIA")
    p_danilo_b = Processo(numero_processo="8888888-88.2026.8.11.0098", empresa_cliente_id=empresa.id, nome_cliente="Y", assistente="DANILO")
    p_danilo_a = Processo(numero_processo="7777777-77.2026.8.11.0097", empresa_cliente_id=empresa.id, nome_cliente="X", assistente="DANILO")
    db.add_all([p_julia, p_danilo_b, p_danilo_a])
    db.flush()
    db.add_all(
        [
            EventoProcesso(processo_id=p_julia.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
            EventoProcesso(processo_id=p_danilo_b.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
            EventoProcesso(processo_id=p_danilo_a.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
        ]
    )
    db.commit()

    relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    numeros = [l.numero_processo for l in relatorio.secoes[0].linhas]
    assert numeros == ["7777777-77.2026.8.11.0097", "8888888-88.2026.8.11.0098", "9999999-99.2026.8.11.0099"]


def test_relatorio_por_empresa_cinco_colunas_e_total_bate(db):
    p1 = _processo(db, numero="1112223-34.2026.8.11.0100", assistente="DANILO")
    p2 = _processo(db, numero="2223334-45.2026.8.11.0101", assistente="JULIA")
    db.add_all(
        [
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS", observacao="obs1"),
            EventoProcesso(processo_id=p2.id, data=date(2026, 9, 5), tipo_evento_nome="DOCUMENTOS"),
        ]
    )
    db.commit()

    relatorio = gerar_relatorio_por_empresa(db, "ABSOLUTA", date(2026, 9, 1), date(2026, 9, 30))
    assert relatorio.total_processos == 2
    assert len(relatorio.linhas) == 2
    linha1 = next(l for l in relatorio.linhas if l.numero_processo == p1.numero_processo)
    assert linha1.assistente == "DANILO"
    assert linha1.cliente == "Fulano de Tal"
    assert linha1.evento == "CUSTAS"
    assert linha1.observacao == "obs1"


def test_relatorio_por_empresa_filtra_por_assistente(db):
    p1 = _processo(db, numero="3334445-56.2026.8.11.0102", assistente="DANILO")
    p2 = _processo(db, numero="4445556-67.2026.8.11.0103", assistente="JULIA")
    db.add_all(
        [
            EventoProcesso(processo_id=p1.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
            EventoProcesso(processo_id=p2.id, data=date(2026, 9, 5), tipo_evento_nome="CUSTAS"),
        ]
    )
    db.commit()

    relatorio = gerar_relatorio_por_empresa(db, "ABSOLUTA", date(2026, 9, 1), date(2026, 9, 30), filtro_assistente="danilo")
    assert relatorio.total_processos == 1
    assert relatorio.linhas[0].assistente == "DANILO"


def test_relatorio_por_empresa_inexistente_gera_erro(db):
    try:
        gerar_relatorio_por_empresa(db, "NAO EXISTE", date(2026, 9, 1), date(2026, 9, 30))
        assert False, "deveria ter levantado ValueError"
    except ValueError as e:
        assert "NAO EXISTE" in str(e)


def test_relatorio_geral_nao_carrega_processos_fora_do_periodo(db):
    """Regressão de performance (já cobria a versão antiga do relatório):
    confirma que o número de consultas SQL não escala com o total de
    processos no banco."""
    dentro = _processo(db, numero="5556667-78.2026.8.11.0104", assistente="DANILO")
    fora = _processo(db, numero="6667778-89.2026.8.11.0105", assistente="JULIA")
    db.add(EventoProcesso(processo_id=dentro.id, data=date(2026, 9, 10), tipo_evento_nome="CUSTAS"))
    db.add(EventoProcesso(processo_id=fora.id, data=date(2020, 1, 1), tipo_evento_nome="CUSTAS"))
    db.commit()

    consultas = []
    engine = db.get_bind()

    def ouvinte(conn, cursor, statement, parameters, context, executemany):
        consultas.append(statement)

    event.listen(engine, "before_cursor_execute", ouvinte)
    try:
        relatorio = gerar_relatorio_geral(db, date(2026, 9, 1), date(2026, 9, 30))
    finally:
        event.remove(engine, "before_cursor_execute", ouvinte)

    assert len(relatorio.secoes) == 1
    assert relatorio.secoes[0].linhas[0].assistente == "DANILO"  # "JULIA" (fora do período) nem aparece
    assert len(consultas) <= 3  # processos do período + último evento por processo
