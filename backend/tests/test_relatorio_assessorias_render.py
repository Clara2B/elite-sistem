"""Fase 7 (Gerador de Relatórios Mensais das Assessorias) — render.py.

Teste de ponta a ponta com dados fictícios (nunca planilhas reais aqui):
o .docx gerado abre, nenhum marcador `{{`/`{%` sobra, e os números/listas
aparecem onde esperado."""
from __future__ import annotations

from datetime import date
from io import BytesIO

from docx import Document

from app.relatorio_assessorias.leitores.base import Origem
from app.relatorio_assessorias.render import (
    DadosRelatorio,
    nome_arquivo_docx,
    renderizar_docx,
)
from app.relatorio_assessorias.secoes.contrarias import (
    ProcessoContrario,
    ResultadoContrarias,
)
from app.relatorio_assessorias.secoes.extrajudiciais import (
    ClienteAusente,
    ResultadoExtrajudiciais,
)
from app.relatorio_assessorias.secoes.iniciais import ResultadoIniciais
from app.relatorio_assessorias.secoes.judiciais import ResultadoAudiencias
from app.relatorio_assessorias.secoes.laudos import ResultadoLaudos
from app.relatorio_assessorias.secoes.manuais import (
    CamposManuais,
    PastasRevisionais,
    SentencaFavoravel,
    SentencaProcedente,
)
from app.relatorio_assessorias.secoes.procons import Procon, ResultadoProcons


def _origem(linha: int) -> Origem:
    return Origem("teste", "aba", linha)


def _dados_fictícios() -> DadosRelatorio:
    return DadosRelatorio(
        assessoria="ASSESSORIA FICTÍCIA",
        mes=8,
        ano=2026,
        data_corte=date(2026, 9, 4),
        laudos=ResultadoLaudos(elaborados=5, entregues_dentro_prazo=4, pendentes=1, pendentes_atrasados=0),
        iniciais=ResultadoIniciais(
            distribuidos_no_mes=3,
            aguardando_distribuicao=1,
            faixas={"Distribuídos em até 7 dias úteis": 2, "Distribuídos entre 8 e 20 dias úteis": 1},
        ),
        extrajudiciais=ResultadoExtrajudiciais(
            enviadas=6,
            realizadas=4,
            pendentes_proximos_meses=2,
            clientes_ausentes=[ClienteAusente("Fulano de Tal", _origem(1))],
        ),
        audiencias_judiciais=ResultadoAudiencias(quantidade=7),
        audiencias_contrarias=ResultadoAudiencias(quantidade=2),
        contrarias=ResultadoContrarias(
            incluidos_no_mes=2,
            ativos_total=3,
            ativos_por_uf={"RJ": 2, "SP": 1},
            lista_judiciais=[
                ProcessoContrario("Beltrano da Silva", "0000001-11.2026.8.19.0001", "RJ", 1500.0, date(2026, 8, 1), _origem(2)),
            ],
            lista_trabalhistas=[
                ProcessoContrario("Ciclana Souza", "0000002-22.2026.5.02.0001", "SP", None, date(2026, 8, 2), _origem(3)),
            ],
        ),
        procons=ResultadoProcons(
            lista=[Procon("Deltrano Costa", "0000003-33.2026.8.13.0001", "EM ANDAMENTO", _origem(4))]
        ),
        manuais=CamposManuais(
            pastas_revisionais=PastasRevisionais(
                recebidas_no_mes=2, aprovadas=1, aguardando_analise=0, aguardando_correcao_no_mes=1, acumulado_crm=4
            ),
            processos_ativos_revisionais_total=2,
            processos_ativos_revisionais_por_uf={"MG": 2},
            sentencas_procedentes=[SentencaProcedente("Elano Pereira", "0000004-44.2026.8.13.0001", "MG")],
            processos_ganhos_por_uf={"MG": 1},
            sentencas_favoraveis_contrarias=[
                SentencaFavoravel("Franjose Lima", "0000005-55.2026.8.19.0001", "RJ", 2000.0)
            ],
            extrajudiciais_solicitacoes_pendentes_correcao=1,
        ),
        avisos=[],
    )


def test_nome_arquivo_docx():
    assert nome_arquivo_docx("EWS", 8, 2026) == "Relatório_EWS_08-2026.docx"


def test_renderizar_docx_abre_sem_marcador_sobrando():
    conteudo = renderizar_docx(_dados_fictícios())
    doc = Document(BytesIO(conteudo))

    texto_total = "\n".join(p.text for p in doc.paragraphs)
    for tabela in doc.tables:
        for linha in tabela.rows:
            texto_total += "\n" + " | ".join(celula.text for celula in linha.cells)

    assert "{{" not in texto_total
    assert "}}" not in texto_total
    assert "{%" not in texto_total


def test_renderizar_docx_preenche_numeros_com_dois_digitos():
    conteudo = renderizar_docx(_dados_fictícios())
    doc = Document(BytesIO(conteudo))
    texto_total = "\n".join(p.text for p in doc.paragraphs)
    assert "JUDICIAIS - 01" in texto_total  # len(lista_judiciais) == 1, com zero à esquerda
    assert "TRABALHISTAS - 01" in texto_total


def test_renderizar_docx_lista_vazia_mostra_so_cabecalho():
    dados = _dados_fictícios()
    dados.procons.lista = []
    conteudo = renderizar_docx(dados)
    doc = Document(BytesIO(conteudo))
    tabela_procons = next(t for t in doc.tables if t.rows[0].cells[0].text == "NOME COMPLETO" and t.rows[0].cells[-1].text == "SITUAÇÃO")
    assert len(tabela_procons.rows) == 1
    texto_total = "\n".join(p.text for p in doc.paragraphs)
    assert "PROCONS - 00" in texto_total


def test_renderizar_docx_duas_imagens_de_mapa():
    conteudo = renderizar_docx(_dados_fictícios())
    doc = Document(BytesIO(conteudo))
    assert len(doc.inline_shapes) == 2


def test_renderizar_docx_valor_causa_nao_informado_vira_texto():
    conteudo = renderizar_docx(_dados_fictícios())
    doc = Document(BytesIO(conteudo))
    texto_total = "\n".join(
        celula.text for tabela in doc.tables for linha in tabela.rows for celula in linha.cells
    )
    assert "Não informado" in texto_total  # lista_trabalhistas[0].valor_causa é None
