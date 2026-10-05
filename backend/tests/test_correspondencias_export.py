"""PDF e Excel de Correspondências (2026-10-05) — conteúdo/formatação.
Testes de rota (auth, Content-Type) ficam em test_api_correspondencias.py,
mesmo padrão dos outros relatórios."""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook

from app.excel_export import gerar_excel_correspondencias
from app.pdf_export import gerar_pdf_correspondencias
from app.services.correspondencias import CorrespondenciasResult, LinhaCorrespondencia

_NOME_LONGO = "Doutora Fulana de Tal Pereira de Albuquerque Nascimento e Silva OAB/SP 123.456"
_FRASE_LONGA = "PROCON / AÇÃO CONTRÁRIA DE REVISÃO DE CONTRATO BANCÁRIO MUITO LONGA PARA TESTAR QUEBRA DE LINHA"


def _resultado(n_linhas=2, valor_invalido=False):
    linhas = []
    for i in range(n_linhas):
        linhas.append(
            LinhaCorrespondencia(
                advogado=_NOME_LONGO, autor=f"Cliente {i}", adv_preposto="ADVOGADO",
                valor_texto="a combinar" if (valor_invalido and i == 0) else "R$ 180,00",
                valor=None if (valor_invalido and i == 0) else 180.0,
                tipo_acao=_FRASE_LONGA,
            )
        )
    total = sum(l.valor for l in linhas if l.valor is not None)
    return CorrespondenciasResult(empresa="Empresa Teste Ltda", mes="JANEIRO", linhas=linhas, total=total)


def test_pdf_gera_arquivo_valido():
    pdf_bytes = gerar_pdf_correspondencias(_resultado())
    assert pdf_bytes[:4] == b"%PDF"


def _contar_paginas(pdf_bytes: bytes) -> int:
    # Conta objetos "/Type /Page" (uma página), não "/Type /Pages" (a
    # árvore-raiz) — sem biblioteca de leitura de PDF disponível neste
    # projeto (não adicionada sem aprovação da Clara), mas o próprio
    # formato interno do PDF já basta pra contar páginas.
    return len(re.findall(rb"/Type\s*/Page[^s]", pdf_bytes))


def test_pdf_com_muitas_linhas_longas_gera_varias_paginas():
    """Confirma que a tabela realmente pagina (reportlab's Table split)
    quando o conteúdo não cabe numa página só — mesmo requisito da Clara de
    repetir o cabeçalho nas páginas seguintes; a verificação visual de que
    o cabeçalho realmente repete foi feita à parte, com uma amostra de PDF
    aberta e aprovada pela Clara."""
    assert _contar_paginas(gerar_pdf_correspondencias(_resultado(n_linhas=2))) == 1
    assert _contar_paginas(gerar_pdf_correspondencias(_resultado(n_linhas=60))) > 1


def test_pdf_nao_quebra_com_valor_invalido():
    pdf_bytes = gerar_pdf_correspondencias(_resultado(valor_invalido=True))
    assert pdf_bytes[:4] == b"%PDF"


def test_excel_tem_cabecalho_mes_e_empresa_e_linhas():
    ws = load_workbook(io.BytesIO(gerar_excel_correspondencias(_resultado()))).active
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][0] == "Empresa: EMPRESA TESTE LTDA"
    assert linhas[1][0] == "Mês: Janeiro"
    assert linhas[3][:5] == ("Advogado", "Autor", "Adv / Preposto", "Valor", "Tipo de ação")
    assert linhas[4][0] == _NOME_LONGO
    assert linhas[4][3] == 180


def test_excel_valor_invalido_vira_texto_nao_numero():
    ws = load_workbook(io.BytesIO(gerar_excel_correspondencias(_resultado(valor_invalido=True)))).active
    linhas = list(ws.iter_rows(values_only=False))
    celula_valor_invalido = linhas[4][3]
    assert celula_valor_invalido.value == "a combinar"
    assert celula_valor_invalido.number_format == "General"


def test_excel_total_soma_so_valores_validos():
    ws = load_workbook(io.BytesIO(gerar_excel_correspondencias(_resultado(n_linhas=3, valor_invalido=True)))).active
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[-1][3:5] == ("Total", 360)  # 2 linhas válidas de R$ 180,00, 1 inválida fora da soma
