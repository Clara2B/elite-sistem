"""Mesmo padrão de tests/test_api_laudos.py — import via API, relatório
JSON, PDF e Excel."""
import io

import openpyxl
from fastapi.testclient import TestClient

from app.db import get_db
from app.main import app


def _planilha_correspondencias_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ADV. CONTRATOS"
    ws.append(["MÊS", "ADVOGADO", "AUTOR", "EMPRESA", "ADV / PREPOSTO", "VALOR", "TIPO DE AÇÃO"])
    ws.append(["Janeiro", "Dra. Fulana", "Cliente A", "EROS", "ADVOGADO", "R$ 180,00", "PROCON"])
    ws.append(["Janeiro", "Dra. Beltrana", "Cliente B", "EROS", "PREPOSTO", "R$ 220,00", "JUDICIAL"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_import_e_relatorio_via_api(db, admin_token):
    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db
    headers = {"Authorization": f"Bearer {admin_token}"}
    try:
        client = TestClient(app)

        resp = client.post(
            "/correspondencias/import",
            files={
                "arquivo": (
                    "correspondencias.xlsx", _planilha_correspondencias_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            headers=headers,
        )
        assert resp.status_code == 200
        assert resp.json()["linhas_novas"] == 2

        resp = client.get("/correspondencias/relatorio", params={"empresa": "EROS", "mes": "Janeiro"}, headers=headers)
        assert resp.status_code == 200
        body = resp.json()
        assert body["total"] == 400.0
        assert len(body["linhas"]) == 2

        resp_pdf = client.get(
            "/correspondencias/relatorio.pdf", params={"empresa": "EROS", "mes": "Janeiro"}, headers=headers
        )
        assert resp_pdf.status_code == 200
        assert resp_pdf.headers["content-type"] == "application/pdf"
        assert resp_pdf.content[:4] == b"%PDF"

        resp_xlsx = client.get(
            "/correspondencias/relatorio.xlsx", params={"empresa": "EROS", "mes": "Janeiro"}, headers=headers
        )
        assert resp_xlsx.status_code == 200
        assert resp_xlsx.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        planilha = openpyxl.load_workbook(io.BytesIO(resp_xlsx.content)).active
        linhas = list(planilha.iter_rows(values_only=True))
        assert linhas[0][0] == "Empresa: EROS"

        resp_sem_login = client.get("/correspondencias/relatorio", params={"empresa": "EROS", "mes": "Janeiro"})
        assert resp_sem_login.status_code == 401
    finally:
        app.dependency_overrides.clear()


def test_relatorio_empresa_nao_encontrada_404(client, db, admin_token):
    resp = client.get(
        "/correspondencias/relatorio", params={"empresa": "NÃO EXISTE", "mes": "Janeiro"},
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert resp.status_code == 404


def test_apagar_todas_correspondencias_via_api_exige_admin(client, db, admin_token):
    resp = client.post(
        "/correspondencias/import",
        files={
            "arquivo": (
                "correspondencias.xlsx", _planilha_correspondencias_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers={"Authorization": f"Bearer {admin_token}"},
    )
    assert resp.status_code == 200

    resp = client.delete("/correspondencias", headers={"Authorization": f"Bearer {admin_token}"})
    assert resp.status_code == 200
    assert resp.json()["apagados"] == 2
