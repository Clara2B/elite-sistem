"""Versão em .xlsx dos relatórios (2026-10-05, a pedido da Clara: "preciso
que todos os relatórios do sistema tenham a opção de serem baixados em
formato de excel"). Testa o conteúdo/formatação das planilhas geradas —
os testes de rota (auth, Content-Disposition) ficam em test_api_laudos.py
e test_web.py, mesmo padrão já usado pros relatórios em PDF."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import load_workbook

from app.auth import hash_senha
from app.excel_export import (
    gerar_excel_audiencias,
    gerar_excel_laudos,
    gerar_excel_processos_geral,
    gerar_excel_processos_por_empresa,
)
from app.models import Audiencia, Usuario
from app.services.audiencias import AudienciasResult
from app.services.empresas import get_or_create_empresa
from app.services.laudos import LaudosResult, LinhaLaudo
from app.services.processos import (
    LinhaProcessoGeral,
    LinhaProcessoPorEmpresa,
    LinhaProcessoResumo,
    RelatorioGeral,
    RelatorioPorEmpresa,
    SecaoEmpresaGeral,
)


def _abrir(xlsx_bytes: bytes):
    assert xlsx_bytes[:2] == b"PK"  # assinatura de um .xlsx (zip) válido
    return load_workbook(io.BytesIO(xlsx_bytes)).active


def test_excel_laudos_linhas_e_formatacao():
    result = LaudosResult(
        empresa="Absoluta", cnpj="00.000.000/0001-00",
        periodo_ini=date(2026, 9, 20), periodo_fim=date(2026, 10, 20), status="SOLICITACAO",
        linhas=[
            LinhaLaudo(data=date(2026, 9, 25), cliente="Fulano", tipo="Laudo X", status="SOLICITACAO", valor=150.0),
            LinhaLaudo(data=date(2026, 9, 26), cliente="Beltrana", tipo="Laudo Y", status="SOLICITACAO", valor=200.0),
        ],
        total=350.0,
    )
    ws = _abrir(gerar_excel_laudos(result))

    cabecalho = [c.value for c in next(r for r in ws.iter_rows() if r[0].value == "Data")]
    assert cabecalho == ["Data", "Cliente", "Tipo de laudo", "Status", "Valor"]

    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][0] == "Empresa: ABSOLUTA"
    assert linhas[1][0] == "CNPJ: 00.000.000/0001-00"
    # uma linha por item, nenhum dado perdido
    assert linhas[5][1:4] == ("Fulano", "Laudo X", "SOLICITACAO")
    assert linhas[6][1:4] == ("Beltrana", "Laudo Y", "SOLICITACAO")
    assert linhas[-1][3:5] == ("Total", 350)

    # data e valor são células de verdade (dá pra somar/filtrar no Excel),
    # não texto formatado à mão.
    celula_data = next(c for r in ws.iter_rows() for c in r if hasattr(c.value, "year") and c.value.year == 2026)
    assert celula_data.number_format == "DD/MM/YYYY"
    celula_valor = next(c for r in ws.iter_rows() for c in r if c.value == 150)
    assert "R$" in celula_valor.number_format

    # o texto "Data" do cabeçalho não pode "herdar" a formatação de data
    celula_cabecalho_data = next(c for r in ws.iter_rows() for c in r if c.value == "Data")
    assert celula_cabecalho_data.number_format == "General"


def test_excel_laudos_sem_cnpj_nao_desalinha_a_tabela():
    result = LaudosResult(
        empresa="Zenith", cnpj=None, periodo_ini=date(2026, 9, 20), periodo_fim=date(2026, 10, 20),
        status="SOLICITACAO",
        linhas=[LinhaLaudo(data=date(2026, 9, 25), cliente="Ciclana", tipo="Laudo Z", status="SOLICITACAO", valor=99.0)],
        total=99.0,
    )
    ws = _abrir(gerar_excel_laudos(result))
    linhas = list(ws.iter_rows(values_only=True))
    # sem CNPJ: Empresa / Status / (branco) / cabeçalho / dado — uma linha a menos que com CNPJ
    assert linhas[0][0] == "Empresa: ZENITH"
    assert linhas[1][0] == "Status: SOLICITACAO"
    assert linhas[3][:5] == ("Data", "Cliente", "Tipo de laudo", "Status", "Valor")
    assert linhas[4][1:4] == ("Ciclana", "Laudo Z", "SOLICITACAO")


def test_excel_audiencias_lista_clientes_numerados():
    result = AudienciasResult(
        empresa="Absoluta", cnpj="00.0", periodo_ini=date(2026, 9, 1), periodo_fim=date(2026, 9, 15),
        valor_unitario=80.0, clientes=["fulano", "beltrana"], total=160.0, quantidade_mes=10,
    )
    ws = _abrir(gerar_excel_audiencias(result))
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][0] == "Empresa: ABSOLUTA"
    assert (1, "FULANO") == linhas[8]
    assert (2, "BELTRANA") == linhas[9]
    assert linhas[-1] == ("Total", 160)


def test_excel_processos_geral_junta_as_duas_partes_por_processo():
    relatorio = RelatorioGeral(
        periodo_ini=date(2026, 9, 1), periodo_fim=date(2026, 9, 30),
        secoes=[
            SecaoEmpresaGeral(
                empresa="Absoluta", total_processos=1,
                linhas=[LinhaProcessoGeral(assistente="Danilo", numero_processo="1234567-12.2026.8.11.0001", evento="CUSTAS", fatal=True)],
                linhas_resumo=[LinhaProcessoResumo(cliente="Fulano", numero_processo="1234567-12.2026.8.11.0001", evento="CUSTAS", observacao="ok")],
            ),
            SecaoEmpresaGeral(
                empresa="Zenith", total_processos=1,
                linhas=[LinhaProcessoGeral(assistente="Ana", numero_processo="7654321-21.2026.8.11.0002", evento="DESPACHO", fatal=False)],
                linhas_resumo=[LinhaProcessoResumo(cliente="Ciclana", numero_processo="7654321-21.2026.8.11.0002", evento="DESPACHO", observacao="pendente")],
            ),
        ],
    )
    ws = _abrir(gerar_excel_processos_geral(relatorio))
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][0] == "Empresa: ELITE MEDIAÇÕES"
    assert linhas[4] == ("Empresa", "Assistente", "Nº Processo", "Cliente", "Evento", "Fatal", "Observação")
    assert linhas[5] == ("Absoluta", "Danilo", "1234567-12.2026.8.11.0001", "Fulano", "CUSTAS", "Sim", "ok")
    assert linhas[6] == ("Zenith", "Ana", "7654321-21.2026.8.11.0002", "Ciclana", "DESPACHO", "Não", "pendente")


def test_excel_processos_por_empresa():
    relatorio = RelatorioPorEmpresa(
        empresa="Absoluta", periodo_ini=date(2026, 9, 1), periodo_fim=date(2026, 9, 30), total_processos=1,
        linhas=[
            LinhaProcessoPorEmpresa(
                assistente="Danilo", numero_processo="1234567-12.2026.8.11.0001",
                cliente="Fulano", evento="CUSTAS", observacao="ok",
            )
        ],
    )
    ws = _abrir(gerar_excel_processos_por_empresa(relatorio))
    linhas = list(ws.iter_rows(values_only=True))
    assert linhas[0][0] == "Empresa: ABSOLUTA"
    assert linhas[4][:5] == ("Assistente", "Nº Processo", "Cliente", "Último evento", "Última observação")
    assert linhas[5][:5] == ("Danilo", "1234567-12.2026.8.11.0001", "Fulano", "CUSTAS", "ok")


# --- Rotas /audiencias/relatorio.xlsx e /processos/relatorio.xlsx -----------
# (/laudos/relatorio.xlsx já é coberto em test_api_laudos.py, mesmo padrão)


def _logar_admin(db, client, email="admin.xlsx@teste.local"):
    db.add(Usuario(nome="Admin", email=email, senha_hash=hash_senha("certa"), papel_global="ADMIN_SUPERIOR"))
    db.commit()
    client.post("/login", data={"email": email, "senha": "certa"})


def test_rota_audiencias_relatorio_xlsx(client, db):
    _logar_admin(db, client, "admin.aud.xlsx@teste.local")
    empresa = get_or_create_empresa(db, "ALFA")
    db.add(Audiencia(empresa_cliente_id=empresa.id, nome_cliente="Fulano", data_recebimento=date(2026, 9, 1)))
    db.commit()

    resposta = client.get(
        "/audiencias/relatorio.xlsx", params={"empresa": "ALFA", "ano": 2026, "mes": 9, "quinzena": 1}
    )
    assert resposta.status_code == 200
    assert resposta.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    planilha = load_workbook(io.BytesIO(resposta.content)).active
    linhas = list(planilha.iter_rows(values_only=True))
    assert linhas[0][0] == "Empresa: ALFA"


def test_rota_processos_relatorio_xlsx_por_empresa(client, db):
    _logar_admin(db, client, "admin.proc.xlsx@teste.local")

    resposta = client.get(
        "/processos/relatorio.xlsx",
        params={"periodo_ini": "2026-01-01", "periodo_fim": "2026-01-31", "tipo": "empresa", "empresa": "ALFA"},
    )
    assert resposta.status_code == 400  # "ALFA" não existe — mesmo comportamento do .pdf


def test_rota_processos_relatorio_xlsx_geral(client, db):
    _logar_admin(db, client, "admin.proc.geral.xlsx@teste.local")

    resposta = client.get(
        "/processos/relatorio.xlsx", params={"periodo_ini": "2026-01-01", "periodo_fim": "2026-01-31", "tipo": "geral"}
    )
    assert resposta.status_code == 200
    assert resposta.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    planilha = load_workbook(io.BytesIO(resposta.content)).active
    linhas = list(planilha.iter_rows(values_only=True))
    assert linhas[4][:7] == ("Empresa", "Assistente", "Nº Processo", "Cliente", "Evento", "Fatal", "Observação")
