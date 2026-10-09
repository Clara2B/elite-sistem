"""Telas /app/relatorios-assessorias (Fase 8) — ponta a ponta via
TestClient, com planilhas pequenas fictícias em memória."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.auth import hash_senha
from app.models import EmpresaCliente, Usuario


def _logar_admin(db, client, email="admin.rel.assessorias@teste.local"):
    db.add(Usuario(nome="Admin", email=email, senha_hash=hash_senha("certa"), papel_global="ADMIN_SUPERIOR"))
    db.commit()
    client.post("/login", data={"email": email, "senha": "certa"})


def _logar_colaborador(db, client, email="colab.rel.assessorias@teste.local"):
    db.add(Usuario(nome="Colaborador", email=email, senha_hash=hash_senha("certa")))
    db.commit()
    client.post("/login", data={"email": email, "senha": "certa"})


def _bytes_planilha(builder) -> bytes:
    wb = builder()
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _laudo():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("AGO26")
    ws.append(["DATA", "EMPRESA", "TIPO DE LAUDO", "ENTRADA DE LAUDO", "LAUDO PRONTO", "DATA", "ENTREGUE A EMPRESA"])
    ws.append(["01/08/2026", "TESTE WEB SA", "AUTO", "", "Sim", "05/08/2026", "Sim"])
    return wb


def _iniciais():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("AGOSTO 2026")
    ws.append(["DATA DE RECEBIMENTO", "ASSESSORIA", "DATA DE DISTRIBUIÇÃO", "PROTOCOLADO", "ESTADOS", "NÚMERO DE PROCESSO"])
    ws.append(["03/08/2026", "TESTE WEB SA", "06/08/2026", "SIM", "RJ", "0001"])
    return wb


def _agendamento():
    wb = Workbook()
    wb.remove(wb.active)
    ws_judicial = wb.create_sheet("JUDICIAL")
    ws_judicial.append([None, "NOME DO CLIENTE", "PROCESSO", "EMPRESA"])
    ws_judicial.append(["01/03/2026", "Fulano", "0001", "TESTE WEB SA"])
    ws_pauta = wb.create_sheet("PAUTA DA SEMANA CONTRARIA")
    ws_pauta.append(["PAUTA DA SEMANA CONTRARIA"])
    ws_pauta.append(["DATA", "NOME DO CLIENTE", "ESTADO", "PROCESSO", "EMPRESA"])
    ws_pauta.append(["01/04/2026", "Beltrano", "SP", "0002", "TESTE WEB SA"])
    return wb


def _extrajudicial():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("AGENDAMENTO")
    ws.append(["DATA DE RECEBIMENTO", "EMPRESA", "NOME COMPLETO", "DATA DE AGENDAMENTO", "PRESENÇA"])
    ws.append(["05/08/2026", "TESTE WEB SA", "Ciclano", "20/08/2026", "PRESENTE"])
    return wb


def _contrarias():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("TESTE WEB SA")
    ws.append(["TESTE WEB SA"])
    ws.append(["título"])
    ws.append(["DATA DE RECEBIMENTO", "NOME DO AUTOR", "Nº DO PROCESSO", "ESTADO", "VALOR DA CAUSA", "STATUS ATUAL", "TIPO AÇÃO"])
    ws.append(["10/08/2026", "Fulano de Tal", "0813055-21.2026.8.10.0001", "RJ", "R$ 1.000,00", "ATIVO", "JUDICIAL"])
    return wb


def _procon():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("TESTE WEB SA")
    ws.append(["TESTE WEB SA"])
    ws.append(["DATA DE RECEBIMENTO", "NOME DO AUTOR", "Nº DO PROCESSO", "ESTADO", "STATUS ATUAL", "TIPO AÇÃO"])
    ws.append(["10/08/2026", "Beltrano", "0002", "SP", "ATIVO", "PROCON"])
    return wb


def _arquivos():
    tipo = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return {
        "laudo": ("laudo.xlsx", _bytes_planilha(_laudo), tipo),
        "iniciais": ("iniciais.xlsx", _bytes_planilha(_iniciais), tipo),
        "agendamento": ("agendamento.xlsx", _bytes_planilha(_agendamento), tipo),
        "extrajudicial": ("extrajudicial.xlsx", _bytes_planilha(_extrajudicial), tipo),
        "contrarias": ("contrarias.xlsx", _bytes_planilha(_contrarias), tipo),
        "procon": ("procon.xlsx", _bytes_planilha(_procon), tipo),
    }


def test_colaborador_sem_acesso(client, db):
    _logar_colaborador(db, client)
    resposta = client.get("/app/relatorios-assessorias", follow_redirects=False)
    assert resposta.status_code == 403


def test_fluxo_completo_upload_painel_revisao_download(client, db):
    _logar_admin(db, client)
    db.add(EmpresaCliente(nome="TESTE WEB SA"))
    db.commit()

    resposta_tela = client.get("/app/relatorios-assessorias")
    assert resposta_tela.status_code == 200

    resposta_upload = client.post(
        "/app/relatorios-assessorias/processar",
        data={"mes": "8", "ano": "2026", "data_corte": "2026-09-04"},
        files=_arquivos(),
        follow_redirects=False,
    )
    assert resposta_upload.status_code == 303
    assert "/painel" in resposta_upload.headers["location"]

    resposta_painel = client.get("/app/relatorios-assessorias/painel", params={"mes": 8, "ano": 2026})
    assert resposta_painel.status_code == 200
    assert "TESTE WEB SA" in resposta_painel.text

    resposta_revisao = client.get(
        "/app/relatorios-assessorias/revisao/TESTE WEB SA", params={"mes": 8, "ano": 2026}
    )
    assert resposta_revisao.status_code == 200
    assert "Avisos" not in resposta_revisao.text or "Números calculados" in resposta_revisao.text

    resposta_sobrescrever = client.post(
        "/app/relatorios-assessorias/revisao/TESTE WEB SA/sobrescrever",
        data={"mes": "8", "ano": "2026", "campo": "laudos.elaborados", "valor": "99"},
        follow_redirects=False,
    )
    assert resposta_sobrescrever.status_code == 303

    resposta_revisao2 = client.get(
        "/app/relatorios-assessorias/revisao/TESTE WEB SA", params={"mes": 8, "ano": 2026}
    )
    assert ">99<" in resposta_revisao2.text

    resposta_manuais = client.post(
        "/app/relatorios-assessorias/revisao/TESTE WEB SA/manuais",
        data={"mes": "8", "ano": "2026", "processos_ativos_revisionais_total": "7"},
        follow_redirects=False,
    )
    assert resposta_manuais.status_code == 303

    resposta_docx = client.get(
        "/app/relatorios-assessorias/docx/TESTE WEB SA", params={"mes": 8, "ano": 2026}
    )
    assert resposta_docx.status_code == 200
    assert resposta_docx.headers["content-type"].startswith("application/vnd.openxmlformats")

    resposta_lote = client.post(
        "/app/relatorios-assessorias/lote",
        data={"mes": "8", "ano": "2026", "assessorias": ["TESTE WEB SA"]},
        follow_redirects=False,
    )
    assert resposta_lote.status_code == 200
    assert resposta_lote.headers["content-type"] == "application/zip"
