"""Fase 2 (Gerador de Relatórios Mensais das Assessorias) — os três
padrões de leitura, localização flexível de cabeçalho, colunas duplicadas
por posição, origem de cada linha. Planilhas pequenas construídas em
memória (sem arquivo binário no repositório)."""
from __future__ import annotations

import pytest
from openpyxl import Workbook

from app.relatorio_assessorias.leitores import (
    aba_mensal,
    aba_por_assessoria,
    aba_unica,
    base,
)

# ---------- base.py: localizar_cabecalho / ler_linhas ----------


def test_localizar_cabecalho_acha_dentro_das_10_primeiras_linhas():
    linhas = [
        ("algum lixo acima",),
        (),
        ("DATA", "EMPRESA", "VALOR"),
        ("01/01/2026", "EWS", 100),
    ]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"], "valor": ["VALOR"]})
    assert localizado is not None
    assert localizado.indice_linha == 2
    assert localizado.mapa_colunas == {"data": 0, "empresa": 1, "valor": 2}


def test_localizar_cabecalho_sem_acento_maiuscula_espaco_extra():
    linhas = [("  dáta  ", "emprêsa")]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"]})
    assert localizado is not None
    assert localizado.mapa_colunas == {"data": 0, "empresa": 1}


def test_localizar_cabecalho_usa_sinonimos():
    linhas = [("UF", "PROCESSO")]
    localizado = base.localizar_cabecalho(linhas, {"estado": ["ESTADO", "UF"], "processo": ["PROCESSO"]})
    assert localizado is not None
    assert localizado.mapa_colunas == {"estado": 0, "processo": 1}


def test_localizar_cabecalho_nao_acha_alem_de_10_linhas():
    linhas = [("lixo",)] * 10 + [("DATA", "EMPRESA")]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"]})
    assert localizado is None


def test_localizar_cabecalho_coluna_faltando_nao_bate():
    linhas = [("DATA", "EMPRESA")]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"], "valor": ["VALOR"]})
    assert localizado is None


def test_colunas_duplicadas_resolvidas_por_posicao():
    # "DATA" aparece duas vezes — a 1ª ocorrência vira "data", a 2ª "data_pronto"
    linhas = [
        ("DATA", "EMPRESA", "DATA"),
        ("01/01/2026", "EWS", "05/01/2026"),
    ]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"], "data_pronto": ["DATA"]})
    assert localizado.mapa_colunas == {"data": 0, "empresa": 1, "data_pronto": 2}
    linhas_lidas = base.ler_linhas(linhas, localizado, "laudo", "AGO26")
    assert len(linhas_lidas) == 1
    assert linhas_lidas[0].valores == {"data": "01/01/2026", "empresa": "EWS", "data_pronto": "05/01/2026"}


def test_coluna_duplicada_sem_segunda_canonica_usa_primeira_ocorrencia():
    # só "empresa" mapeado pra "EMPRESA" — com EMPRESA duas vezes, usa a 1ª
    linhas = [("EMPRESA", "PROCESSO", "EMPRESA"), ("EWS", "123", "SW")]
    localizado = base.localizar_cabecalho(linhas, {"empresa": ["EMPRESA"], "processo": ["PROCESSO"]})
    assert localizado.mapa_colunas == {"empresa": 0, "processo": 1}


def test_colunas_fixas_nao_precisa_de_titulo():
    # coluna A sem título (JUDICIAL) — "data" fixada no índice 0
    linhas = [
        (None, "NOME DO CLIENTE", "PROCESSO"),
        ("01/01/2026", "Fulano", "123"),
    ]
    localizado = base.localizar_cabecalho(
        linhas, {"nome_cliente": ["NOME DO CLIENTE"], "processo": ["PROCESSO"]}, colunas_fixas={"data": 0}
    )
    assert localizado is not None
    assert localizado.mapa_colunas == {"nome_cliente": 1, "processo": 2, "data": 0}


def test_ler_linhas_descarta_linha_vazia():
    linhas = [("DATA", "EMPRESA"), ("01/01/2026", "EWS"), (None, None), ("", "")]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"]})
    linhas_lidas = base.ler_linhas(linhas, localizado, "laudo", "AGO26")
    assert len(linhas_lidas) == 1


def test_ler_linhas_descarta_titulo_repetido_no_meio():
    linhas = [
        ("DATA", "EMPRESA"),
        ("01/01/2026", "EWS"),
        ("DATA", "EMPRESA"),  # título repetido, não é dado
        ("02/01/2026", "SW"),
    ]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"]})
    linhas_lidas = base.ler_linhas(linhas, localizado, "laudo", "AGO26")
    assert len(linhas_lidas) == 2
    assert [l.valores["empresa"] for l in linhas_lidas] == ["EWS", "SW"]


def test_origem_de_cada_linha():
    linhas = [("DATA", "EMPRESA"), ("01/01/2026", "EWS"), ("02/01/2026", "SW")]
    localizado = base.localizar_cabecalho(linhas, {"data": ["DATA"], "empresa": ["EMPRESA"]})
    linhas_lidas = base.ler_linhas(linhas, localizado, "laudo", "AGO26")
    assert linhas_lidas[0].origem == base.Origem("laudo", "AGO26", 2)
    assert linhas_lidas[1].origem == base.Origem("laudo", "AGO26", 3)


def test_ler_aba_levanta_erro_bloqueante_sem_coluna():
    linhas = [("DATA",)]
    with pytest.raises(base.CabecalhoNaoEncontrado):
        base.ler_aba(linhas, {"colunas": {"data": ["DATA"], "empresa": ["EMPRESA"]}}, "laudo", "AGO26")


# ---------- aba_mensal.py ----------


@pytest.mark.parametrize(
    "nome_aba",
    ["AGO26", "ago26", "AGO 26", "AGOSTO 2026", "agosto2026", "AGOSTO26"],
)
def test_encontrar_aba_do_mes_formatos_flexiveis(nome_aba):
    assert aba_mensal.encontrar_aba_do_mes([nome_aba], mes=8, ano=2026) == nome_aba


def test_encontrar_aba_do_mes_ignora_acento():
    assert aba_mensal.encontrar_aba_do_mes(["MARCO 2026"], mes=3, ano=2026) == "MARCO 2026"
    assert aba_mensal.encontrar_aba_do_mes(["MARÇO 2026"], mes=3, ano=2026) == "MARÇO 2026"


def test_encontrar_aba_do_mes_nao_acha():
    assert aba_mensal.encontrar_aba_do_mes(["JAN26", "FEV26"], mes=8, ano=2026) is None


def test_encontrar_todas_abas_do_ano():
    abas = ["JAN26", "FEV26", "AGO26", "DASHBOARD", "DADOS"]
    resultado = aba_mensal.encontrar_todas_abas_do_ano(abas, 2026)
    assert resultado == {1: "JAN26", 2: "FEV26", 8: "AGO26"}


def _workbook_laudos_agosto():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("AGO26")
    ws.append(["DATA", "EMPRESA", "TIPO DE LAUDO", "ENTRADA DE LAUDO", "LAUDO PRONTO", "DATA", "ENTREGUE A EMPRESA"])
    ws.append(["01/08/2026", "EWS", "AUTO", "Normal", "Sim", "05/08/2026", "Sim"])
    wb.create_sheet("DASHBOARD")
    return wb


def test_aba_mensal_ler_end_to_end():
    wb = _workbook_laudos_agosto()
    config = {
        "colunas": {
            "data": ["DATA"], "empresa": ["EMPRESA"], "tipo_laudo": ["TIPO DE LAUDO"],
            "entrada_laudo": ["ENTRADA DE LAUDO"], "laudo_pronto": ["LAUDO PRONTO"],
            "data_pronto": ["DATA"], "entregue_a_empresa": ["ENTREGUE A EMPRESA"],
        }
    }
    linhas = aba_mensal.ler(wb, "laudo", config, mes=8, ano=2026)
    assert len(linhas) == 1
    assert linhas[0].valores["empresa"] == "EWS"
    assert linhas[0].valores["data_pronto"] == "05/08/2026"
    assert linhas[0].origem.aba == "AGO26"


def test_aba_mensal_ler_levanta_aba_nao_encontrada():
    wb = _workbook_laudos_agosto()
    with pytest.raises(base.AbaNaoEncontrada):
        aba_mensal.ler(wb, "laudo", {"colunas": {}}, mes=12, ano=2026)


# ---------- aba_unica.py ----------


def test_aba_unica_ler_end_to_end():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("JUDICIAL")
    ws.append([None, "NOME DO CLIENTE", "PROCESSO", "EMPRESA"])
    ws.append(["01/08/2026", "Fulano de Tal", "0001", "EWS"])
    config = {
        "aba": "JUDICIAL",
        "colunas": {"nome_cliente": ["NOME DO CLIENTE"], "processo": ["PROCESSO"], "empresa": ["EMPRESA"]},
        "colunas_fixas": {"data": 0},
    }
    linhas = aba_unica.ler(wb, "agendamento", config)
    assert len(linhas) == 1
    assert linhas[0].valores == {"nome_cliente": "Fulano de Tal", "processo": "0001", "empresa": "EWS", "data": "01/08/2026"}


def test_aba_unica_cabecalho_fora_da_linha_1():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("PAUTA DA SEMANA CONTRARIA")
    ws.append(["PAUTA DA SEMANA CONTRARIA"])  # linha 1: só título
    ws.append(["DATA", "NOME DO CLIENTE", "ESTADO", "PROCESSO", "EMPRESA"])
    ws.append(["01/08/2026", "Fulano", "SP", "0001", "EWS"])
    config = {
        "aba": "PAUTA DA SEMANA CONTRARIA",
        "colunas": {
            "data": ["DATA"], "nome_cliente": ["NOME DO CLIENTE"], "estado": ["ESTADO"],
            "processo": ["PROCESSO"], "empresa": ["EMPRESA"],
        },
    }
    linhas = aba_unica.ler(wb, "agendamento", config)
    assert len(linhas) == 1
    assert linhas[0].origem.linha == 3


def test_aba_unica_levanta_aba_nao_encontrada():
    wb = Workbook()
    with pytest.raises(base.AbaNaoEncontrada):
        aba_unica.ler(wb, "agendamento", {"aba": "JUDICIAL", "colunas": {}})


# ---------- aba_por_assessoria.py ----------


def _workbook_contrarias():
    wb = Workbook()
    wb.remove(wb.active)
    for nome_aba, autor in [("EWS", "Fulano"), ("SW", "Beltrano"), ("NOVARE", "Ciclano")]:
        ws = wb.create_sheet(nome_aba)
        ws.append([nome_aba])
        ws.append(["título qualquer"])
        ws.append(["DATA DE RECEBIMENTO", "NOME DO AUTOR", "Nº DO PROCESSO", "ESTADO", "STATUS ATUAL", "TIPO AÇÃO"])
        ws.append(["01/08/2026", autor, "0001", "SP", "ATIVO", "JUDICIAL"])
    wb.create_sheet("DASHBOARD")
    return wb


_CONFIG_CONTRARIAS = {
    "colunas": {
        "data_recebimento": ["DATA DE RECEBIMENTO"], "nome_autor": ["NOME DO AUTOR"],
        "numero_processo": ["Nº DO PROCESSO"], "estado": ["ESTADO"],
        "status_atual": ["STATUS ATUAL"], "tipo_acao": ["TIPO AÇÃO"],
    }
}


def test_aba_por_assessoria_soma_apelidos():
    wb = _workbook_contrarias()
    nomes_por_apelido = {"EWS": "EWS", "SW": "EWS", "NOVARE": "NOVARE"}
    resultado = aba_por_assessoria.ler(wb, "contrarias", _CONFIG_CONTRARIAS, nomes_por_apelido)
    assert len(resultado["EWS"]) == 2  # abas EWS + SW somadas
    assert [l.valores["nome_autor"] for l in resultado["EWS"]] == ["Fulano", "Beltrano"]
    assert len(resultado["NOVARE"]) == 1


def test_aba_por_assessoria_ignora_aba_sem_apelido_conhecido():
    wb = _workbook_contrarias()
    nomes_por_apelido = {"EWS": "EWS"}  # SW e NOVARE não cadastrados
    resultado = aba_por_assessoria.ler(wb, "contrarias", _CONFIG_CONTRARIAS, nomes_por_apelido)
    assert set(resultado.keys()) == {"EWS"}
    assert len(resultado["EWS"]) == 1  # só a aba "EWS", não "SW"


def test_aba_por_assessoria_aba_malformada_levanta_erro():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("EWS")
    ws.append(["coluna errada"])
    with pytest.raises(base.CabecalhoNaoEncontrado):
        aba_por_assessoria.ler(wb, "contrarias", _CONFIG_CONTRARIAS, {"EWS": "EWS"})
