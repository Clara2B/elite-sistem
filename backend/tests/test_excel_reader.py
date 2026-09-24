import tempfile
from datetime import date
from pathlib import Path

import openpyxl

from app.excel_reader import _normalizar_larguras, load_data_sheets
from app.utils import parse_date_cell


def test_normalizar_larguras_preenche_linhas_curtas_com_none():
    rows = [
        ("A", "B", "C"),
        ("1", "2", "3", "4", "5", "6", "7", "8", "9"),
        ("x", "y", "z"),
    ]
    resultado = _normalizar_larguras(rows)
    assert all(len(r) == 9 for r in resultado)
    assert resultado[0] == ("A", "B", "C", None, None, None, None, None, None)
    assert resultado[1] == ("1", "2", "3", "4", "5", "6", "7", "8", "9")
    assert resultado[2] == ("x", "y", "z", None, None, None, None, None, None)


def test_normalizar_larguras_lista_vazia():
    assert _normalizar_larguras([]) == []


def test_normalizar_larguras_ja_uniforme_fica_igual():
    rows = [("A", "B"), ("C", "D")]
    assert _normalizar_larguras(rows) == rows


def test_load_data_sheets_preenche_col_a_pela_posicao_nao_pelo_nome():
    """`_COL_A` guarda o valor bruto da coluna A de cada linha, por posição
    — não por nome de cabeçalho. Usado pelo import de laudos porque a
    planilha real tem abas onde mais de uma coluna cai no mesmo nome
    canônico "DATA" (a ordem das colunas varia de aba pra aba), fazendo a
    busca por nome pegar a coluna errada em algumas abas mas não em
    outras — a Clara confirmou que a data certa é sempre a da coluna A,
    independente do que o nome do cabeçalho diz. Ver DECISIONS.md."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AGOSTO"
    ws.append(["DATA", "EMPRESA", "TIPO DE LAUDO", "NOME DO CLIENTE"])
    ws.append([date(2026, 8, 26), "ABSOLUTA", "EMPRÉSTIMO", "JOSE NUNES DA ROSA FILHO"])
    ws2 = wb.create_sheet("SETEMBRO")
    # Aba com layout diferente: coluna A continua sendo a data certa, mas
    # a ordem das outras colunas mudou em relação à aba anterior — o que
    # já bastou, na planilha real, pra confundir a busca por nome em algum
    # ponto do meio.
    ws2.append(["DATA", "TIPO DE LAUDO", "EMPRESA", "NOME DO CLIENTE"])
    ws2.append([date(2026, 9, 11), "AUTO", "ABSOLUTA", "VILMAR MARQUES"])
    path = Path(tempfile.mkdtemp()) / "laudos.xlsx"
    wb.save(path)

    df = load_data_sheets(str(path), ["EMPRESA", "TIPO DE LAUDO", "DATA"])
    assert "_COL_A" in df.columns
    valores = sorted(parse_date_cell(v) for v in df["_COL_A"])
    assert valores == [date(2026, 8, 26), date(2026, 9, 11)]


def test_load_data_sheets_marca_data_e_liberacao_so_na_aba_que_usa_esse_alias():
    """`_DATA_E_LIBERACAO` sinaliza, por aba, quando o que virou "DATA" (via
    HEADER_ALIASES) veio da coluna de data de liberação/intake — não de uma
    data de andamento real (ver app/services/processos.py). Uma aba com
    "DIA" (data real) não deve ser marcada; só a que usa literalmente
    "DATA DE LIBERAÇÃO - QUANDO A DRA INSERIIU O CLIENTE NA PLANILHA"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FATAIS 08"
    ws.append(["CLIENTE", "Nº PROCESSO", "DIA"])
    ws.append(["Fulano de Tal", "1234567-11.2026.8.11.0001", date(2026, 9, 5)])
    ws2 = wb.create_sheet("Dra Teste")
    ws2.append(["CLIENTE", "Nº PROCESSO", "DATA DE LIBERAÇÃO - QUANDO A DRA INSERIIU O CLIENTE NA PLANILHA"])
    ws2.append(["Beltrano", "7654321-22.2026.8.11.0002", date(2026, 9, 6)])
    path = Path(tempfile.mkdtemp()) / "processos.xlsx"
    wb.save(path)

    df = load_data_sheets(str(path), ["CLIENTE", "Nº PROCESSO"])
    marcadas = dict(zip(df["_ABA"], df["_DATA_E_LIBERACAO"], strict=True))
    assert marcadas == {"FATAIS 08": False, "Dra Teste": True}
