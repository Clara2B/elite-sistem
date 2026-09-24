"""Leitura genérica das planilhas .xlsx, isolando as abas de dados reais das
abas de controle/dashboard, e unificando cabeçalhos equivalentes entre abas
diferentes.

Portado de core/excel_reader.py do sistema atual (leitor-relatorio) sem
mudança de comportamento.
"""
from __future__ import annotations

import itertools
import logging
import time

import openpyxl
import pandas as pd

from app.utils import normalize

_logger = logging.getLogger("elite_sistem.import")

HEADER_ALIASES = {
    normalize("ENTRADA DO LAUDO"): normalize("ENTRADA DE LAUDO"),
    normalize("VALOR FALTANTE"): normalize("VALOR"),
    normalize("PAGO (SIM/NÃO)"): normalize("PAGO"),
    # Planilha "ELITE - GESTÃO DE PROCESSOS.xlsx" (Fase 5): mesmo cabeçalho
    # de fato aparece com textos diferentes entre abas (mensais, por
    # advogada, "fatais") — unifica pro parser tratar tudo igual.
    normalize("DIA"): normalize("DATA"),
    normalize("DATA DE LIBERAÇÃO - QUANDO A DRA INSERIIU O CLIENTE NA PLANILHA"): normalize("DATA"),
    normalize("FATALISSIMO"): normalize("PRAZO FATAL"),
    normalize("FATAL"): normalize("PRAZO FATAL"),
    normalize("DRA"): normalize("ADVOGADA"),
    normalize("EVENTO - DRA IQUE INSERI AS INFORMAÇÕES A SEREM FEITAS PARA O CLIENTE"): normalize("EVENTO"),
    normalize(
        "OBSERVAÇÃO - INFORMAR O QUE FOI REALIZADO OU O QUE ESTA EM ABERTP PARA ESSE CLIENTE"
    ): normalize("OBSERVAÇÃO"),
    normalize(
        "OBSERVAÇÃO - INFORMAR O QUE FOI REALIZADO OU O QUE ESTA EM ABERTO PARA ESSE CLIENTE"
    ): normalize("OBSERVAÇÃO"),
}


def _canonical_header(texto: str) -> str:
    chave = normalize(texto)
    return HEADER_ALIASES.get(chave, chave)


# Texto exato (antes do alias) que HEADER_ALIASES equipara a "DATA" — mas
# semanticamente é outra coisa: a data em que a Dra inseriu o cliente na
# planilha (intake), não a data do andamento em si. Algumas abas "coringa"
# da planilha real de Gestão de Processos (fatais, Dra Galzo, DOCS E CUSTAS
# etc.) só têm essa coluna, sem nenhuma data de andamento de verdade — ver
# app/services/processos.py, que usa esse marcador pra excluir essas linhas
# do filtro por período do relatório mensal (a Clara confirmou: não deve
# contar como se fosse o mês do andamento).
_TEXTO_DATA_DE_LIBERACAO = normalize(
    "DATA DE LIBERAÇÃO - QUANDO A DRA INSERIIU O CLIENTE NA PLANILHA"
)


def _find_header_row(rows: list[tuple], required_headers: list[str]) -> int | None:
    required_canonicos = [_canonical_header(h) for h in required_headers]
    for i, row in enumerate(rows[:5]):
        values = [_canonical_header(c) for c in row if c is not None]
        if all(h in values for h in required_canonicos):
            return i
    return None


def _normalizar_larguras(rows: list[tuple]) -> list[tuple]:
    """Em modo `read_only=True` do openpyxl, uma linha sem célula preenchida
    no fim pode vir "cortada" (tupla mais curta que outras linhas da mesma
    aba) — o comprimento de cada linha reflete só as células realmente
    escritas naquele trecho do XML, diferente do modo padrão (que sempre
    preenche até a última coluna usada na aba inteira). Sem isso, o pandas
    recusa montar o DataFrame quando linhas têm tamanhos diferentes
    ("X columns passed, passed data had Y columns")."""
    if not rows:
        return rows
    largura = max(len(r) for r in rows)
    return [r + (None,) * (largura - len(r)) for r in rows]


def load_data_sheets(
    path: str,
    required_headers: list[str],
    chave_duplicidade: list[str] | None = None,
) -> pd.DataFrame:
    """Lê todas as abas do arquivo que contenham as colunas exigidas e devolve
    tudo empilhado num único DataFrame, com uma coluna extra '_ABA'.

    `read_only=True`: o modo padrão do openpyxl carrega a planilha inteira
    como objetos Python na memória, o que fica pesado em planilhas com
    dezenas de milhares de linhas (ex.: Gestão de Processos, Fase 5) —
    plausível causa de estourar a memória do plano gratuito do Render. Modo
    somente leitura lê linha a linha, sem esse custo; só leitura de valor
    de célula (`iter_rows(values_only=True)`) é usada aqui, então não perde
    nada."""
    inicio = time.perf_counter()
    abas_aproveitadas = 0
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        frames = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            iterador = ws.iter_rows(values_only=True)
            # Só as 5 primeiras linhas (o quanto `_find_header_row` olha)
            # pra decidir se essa aba interessa, antes de puxar o resto do
            # arquivo pra memória. Uma planilha real tem ~21 abas (mensais,
            # por advogada, "fatais", dashboard) e a maioria não bate com os
            # cabeçalhos exigidos — sem esse corte, cada uma delas era lida
            # (e transformada em linha Python) por inteiro só pra ser
            # descartada logo depois, o que pesa bastante numa aba grande
            # que acaba não servindo pra nada.
            primeiras_linhas = list(itertools.islice(iterador, 5))
            if not primeiras_linhas:
                continue
            header_idx = _find_header_row(primeiras_linhas, required_headers)
            if header_idx is None:
                continue
            rows = _normalizar_larguras(primeiras_linhas + list(iterador))
            header = [
                _canonical_header(c) if c is not None else f"col_{i}"
                for i, c in enumerate(rows[header_idx])
            ]
            seen: dict[str, int] = {}
            cols = []
            for h in header:
                seen[h] = seen.get(h, 0) + 1
                cols.append(h if seen[h] == 1 else f"{h}_{seen[h]}")
            data_rows = rows[header_idx + 1:]
            df = pd.DataFrame(data_rows, columns=cols)
            df["_ABA"] = sheet_name
            # Valor bruto da coluna A (posição, não nome) — reserva pra quando
            # o nome do cabeçalho não é confiável: a planilha real de Laudos
            # tem abas com mais de uma coluna que cai no mesmo nome canônico
            # "DATA" (ex.: um alias como "DIA" batendo com outra coluna de
            # data, tipo prazo/entrega, dependendo da ordem das colunas
            # naquela aba específica) — o nome resolve pra coluna errada só
            # nalgumas abas, não em todas. A Clara confirmou: a data certa é
            # sempre a da coluna A. Ver app/services/laudos.py.
            df["_COL_A"] = [r[0] if r else None for r in data_rows]
            # True quando o que essa aba chama de "DATA" (depois do alias em
            # HEADER_ALIASES) na verdade veio da coluna de data de liberação/
            # intake, não de uma data de andamento real — ver comentário de
            # `_TEXTO_DATA_DE_LIBERACAO` acima.
            df["_DATA_E_LIBERACAO"] = any(
                c is not None and normalize(c) == _TEXTO_DATA_DE_LIBERACAO for c in rows[header_idx]
            )
            df = df.dropna(how="all", subset=[c for c in cols if c != "_ABA"])
            frames.append(df)
            abas_aproveitadas += 1
    finally:
        wb.close()  # modo read_only mantém o arquivo aberto até fechar explicitamente
    if not frames:
        _logger.info(
            "load_data_sheets: %.1fs, nenhuma aba com os cabeçalhos exigidos (de %d abas no arquivo)",
            time.perf_counter() - inicio, len(wb.sheetnames),
        )
        return pd.DataFrame()
    resultado = pd.concat(frames, ignore_index=True)
    if chave_duplicidade:
        colunas_chave = [_canonical_header(h) for h in chave_duplicidade]
        colunas_chave = [c for c in colunas_chave if c in resultado.columns]
    else:
        colunas_chave = []
    if not colunas_chave:
        colunas_chave = [c for c in resultado.columns if c != "_ABA"]
    resultado = resultado.drop_duplicates(subset=colunas_chave, keep="first")
    resultado = resultado.reset_index(drop=True)
    # Diagnóstico pra próxima vez que "ler a planilha" for reportado como
    # lento (ver DECISIONS.md) — sem isso só dava pra adivinhar se o tempo
    # estava indo no parse do arquivo ou no resto do import (gravar no
    # banco).
    _logger.info(
        "load_data_sheets: %.1fs, %d linha(s) em %d/%d aba(s) aproveitada(s)",
        time.perf_counter() - inicio, len(resultado), abas_aproveitadas, len(wb.sheetnames),
    )
    return resultado
