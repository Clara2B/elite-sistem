"""Versão em .xlsx dos relatórios que já existem em PDF (2026-10-05, a
pedido da Clara: "preciso que todos os relatórios do sistema tenham a
opção de serem baixados em formato de excel"). Mesmos dados de
`app/pdf_export.py`, reorganizados em tabela "plana" (uma linha por item,
sem quebra de página/seção visual) — o formato que facilita filtrar/somar
no próprio Excel, em vez de recriar o layout do PDF célula por célula.

Cartas fica de fora: não é um relatório tabular (é texto de carta-convite
em PDF, que já ganhou "Copiar texto" — ver services/cartas.py). Pendências
também fica de fora: a tela de Pendências gera uma mensagem de cobrança em
texto corrido, não uma lista/tabela."""
from __future__ import annotations

import io
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

if TYPE_CHECKING:
    from app.services.audiencias import AudienciasResult
    from app.services.laudos import LaudosResult
    from app.services.processos import RelatorioGeral, RelatorioPorEmpresa

_NEGRITO = Font(bold=True)
_FORMATO_DATA = "DD/MM/YYYY"
_FORMATO_MOEDA = '"R$" #,##0.00'


def _autosize(ws: Worksheet, linhas: list[list]) -> None:
    larguras: dict[int, int] = {}
    for linha in linhas:
        for i, valor in enumerate(linha, start=1):
            larguras[i] = max(larguras.get(i, 0), len(str(valor)) if valor is not None else 0)
    for i, largura in larguras.items():
        ws.column_dimensions[get_column_letter(i)].width = min(max(largura + 2, 10), 60)


def _cabecalho_tabela(ws: Worksheet, colunas: list[str]) -> None:
    ws.append(colunas)
    for cel in ws[ws.max_row]:
        cel.font = _NEGRITO


def _para_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_excel_laudos(result: LaudosResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Laudos"

    ws.append(["Empresa", result.empresa])
    ws["A1"].font = _NEGRITO
    if result.cnpj:
        ws.append(["CNPJ", result.cnpj])
        ws[f"A{ws.max_row}"].font = _NEGRITO
    ws.append(["Status", result.status])
    ws[f"A{ws.max_row}"].font = _NEGRITO
    ws.append([])

    _cabecalho_tabela(ws, ["Data", "Cliente", "Tipo de laudo", "Status", "Valor"])
    # Captura a linha do cabeçalho DEPOIS de escrevê-lo, nunca antes: um
    # `ws.append([])` (linha em branco, usada acima como separador) não
    # necessariamente avança `ws.max_row` do jeito previsível — calcular a
    # posição de antemão (`max_row + 1`) pode sair errado por uma linha.
    linha_cabecalho = ws.max_row
    linhas_dados = [[linha.data, linha.cliente, linha.tipo, linha.status, linha.valor] for linha in result.linhas]
    for linha in linhas_dados:
        ws.append(linha)
    ws.append(["", "", "", "Total", result.total])
    ws[f"D{ws.max_row}"].font = _NEGRITO
    ws[f"E{ws.max_row}"].font = _NEGRITO

    for row in ws.iter_rows(min_row=linha_cabecalho + 1, max_row=linha_cabecalho + len(linhas_dados), min_col=1, max_col=1):
        row[0].number_format = _FORMATO_DATA
    for row in ws.iter_rows(min_row=linha_cabecalho + 1, max_row=ws.max_row, min_col=5, max_col=5):
        row[0].number_format = _FORMATO_MOEDA

    _autosize(ws, [[result.empresa]] + linhas_dados + [["Tipo de laudo", "Cliente"]])
    return _para_bytes(wb)


def gerar_excel_audiencias(result: AudienciasResult) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Audiências"

    ws.append(["Empresa", result.empresa])
    ws["A1"].font = _NEGRITO
    if result.cnpj:
        ws.append(["CNPJ", result.cnpj])
        ws[f"A{ws.max_row}"].font = _NEGRITO
    ws.append(["Período", f"{result.periodo_ini.strftime('%d/%m/%Y')} - {result.periodo_fim.strftime('%d/%m/%Y')}"])
    ws[f"A{ws.max_row}"].font = _NEGRITO
    if result.valor_unitario is not None:
        ws.append(["Valor por audiência", result.valor_unitario])
        ws[f"A{ws.max_row}"].font = _NEGRITO
        ws[f"B{ws.max_row}"].number_format = _FORMATO_MOEDA
    ws.append(["Quantidade solicitada no período", len(result.clientes)])
    ws[f"A{ws.max_row}"].font = _NEGRITO
    ws.append(["Acumulado no mês", result.quantidade_mes])
    ws[f"A{ws.max_row}"].font = _NEGRITO
    ws.append([])

    _cabecalho_tabela(ws, ["#", "Cliente"])
    for i, nome in enumerate(result.clientes, start=1):
        ws.append([i, nome.upper()])

    if result.total is not None:
        ws.append(["", "Total"])
        ws.append(["", result.total])
        ws[f"B{ws.max_row - 1}"].font = _NEGRITO
        ws[f"B{ws.max_row}"].number_format = _FORMATO_MOEDA

    _autosize(ws, [[result.empresa]] + [[i, n.upper()] for i, n in enumerate(result.clientes, start=1)])
    return _para_bytes(wb)


def gerar_excel_processos_geral(relatorio: RelatorioGeral) -> bytes:
    """Uma linha por processo, com Empresa como coluna — junta as duas
    partes que o PDF mostra separadas (Parte 1: Assistente/Nº processo/
    Evento/Fatal; Parte 2: Cliente/Nº processo/Último evento/Última
    observação, ver pdf_export.py::gerar_pdf_processos_geral) numa tabela
    só, casando pelo nº do processo — mais fácil de filtrar/somar no Excel
    do que duas tabelas repetindo o mesmo processo."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Processos - Geral"

    periodo = f"{relatorio.periodo_ini.strftime('%d/%m/%Y')} a {relatorio.periodo_fim.strftime('%d/%m/%Y')}"
    ws.append(["Período", periodo])
    ws["A1"].font = _NEGRITO
    ws.append([])

    _cabecalho_tabela(ws, ["Empresa", "Assistente", "Nº Processo", "Cliente", "Evento", "Fatal", "Observação"])
    linhas_dados = []
    for secao in relatorio.secoes:
        resumo_por_numero = {l.numero_processo: l for l in secao.linhas_resumo}
        for l in secao.linhas:
            resumo = resumo_por_numero.get(l.numero_processo)
            linhas_dados.append([
                secao.empresa, l.assistente, l.numero_processo,
                resumo.cliente if resumo else "—", l.evento, "Sim" if l.fatal else "Não",
                resumo.observacao if resumo else "—",
            ])
    for linha in linhas_dados:
        ws.append(linha)

    _autosize(ws, [["Empresa", "Assistente", "Nº Processo", "Cliente", "Evento", "Fatal", "Observação"]] + linhas_dados)
    return _para_bytes(wb)


def gerar_excel_processos_por_empresa(relatorio: RelatorioPorEmpresa) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Processos"

    periodo = f"{relatorio.periodo_ini.strftime('%d/%m/%Y')} a {relatorio.periodo_fim.strftime('%d/%m/%Y')}"
    ws.append(["Empresa", relatorio.empresa])
    ws["A1"].font = _NEGRITO
    ws.append(["Período", periodo])
    ws[f"A{ws.max_row}"].font = _NEGRITO
    ws.append(["Total de processos", relatorio.total_processos])
    ws[f"A{ws.max_row}"].font = _NEGRITO
    ws.append([])

    _cabecalho_tabela(ws, ["Assistente", "Nº Processo", "Cliente", "Último evento", "Última observação"])
    linhas_dados = [
        [l.assistente, l.numero_processo, l.cliente, l.evento, l.observacao] for l in relatorio.linhas
    ]
    for linha in linhas_dados:
        ws.append(linha)

    _autosize(ws, [[relatorio.empresa], ["Assistente", "Nº Processo", "Cliente", "Último evento"]] + linhas_dados)
    return _para_bytes(wb)
