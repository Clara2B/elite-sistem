"""Orquestra leitura + normalização + cálculo + validação pras 6 planilhas
de um mês, pra TODAS as assessorias do cadastro de uma vez (Fase 8,
especificação: "O sistema lê as 6 fontes, normaliza e calcula os dados de
todas as assessorias do cadastro").

Cada fonte é lida UMA VEZ (não uma vez por assessoria) e depois filtrada
por assessoria (`normalizacao.assessoria.filtrar_linhas` pras fontes de
`aba_mensal`/`aba_unica`; Contrárias e Procon já vêm agrupadas por
`aba_por_assessoria.ler`). Erros bloqueantes (`leitores.base.
ErroDeLeituraBloqueante` — arquivo errado, aba obrigatória ausente, coluna
obrigatória não encontrada) sobem direto pra quem chamou (a rota de
upload), que mostra o erro e não processa nada — não faz sentido salvar
resultado parcial se uma das 6 fontes não leu."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import EmpresaCliente, Usuario
from app.relatorio_assessorias import armazenamento, validacao
from app.relatorio_assessorias.avisos import Aviso
from app.relatorio_assessorias.config import carregar
from app.relatorio_assessorias.leitores import aba_mensal, aba_por_assessoria, aba_unica
from app.relatorio_assessorias.models import RelatorioAssessoriaResultado
from app.relatorio_assessorias.normalizacao.assessoria import (
    construir_nomes_por_apelido,
    filtrar_linhas,
)
from app.relatorio_assessorias.render import DadosRelatorio
from app.relatorio_assessorias.secoes import (
    contrarias,
    extrajudiciais,
    iniciais,
    judiciais,
    laudos,
    manuais,
    procons,
)

# Chaves == `arquivo_campo_upload` de cada fonte em `config/fontes.yaml` —
# as 6 planilhas dos "Seis campos fixos" da especificação ("Fluxo do
# usuário e telas"). "agendamento" serve duas fontes (Judicial e Pauta da
# Semana Contrária, mesmo arquivo "PLANILHA 2026").
CAMPOS_UPLOAD = ["laudo", "iniciais", "agendamento", "extrajudicial", "contrarias", "procon"]


@dataclass
class EmpresaDesconhecidaGlobal:
    nome: str
    quantidade: int


def processar_upload(
    db: Session,
    caminhos: dict[str, str],
    mes: int,
    ano: int,
    data_corte: date,
    usuario: Usuario,
) -> tuple[list[RelatorioAssessoriaResultado], list[EmpresaDesconhecidaGlobal]]:
    """`caminhos`: {campo_upload: caminho do arquivo salvo em disco (temp)}
    — ver `CAMPOS_UPLOAD`. Devolve os resultados salvos (um por assessoria
    ativa do cadastro) e a lista de nomes de empresa encontrados nas
    planilhas que não batem com nenhuma assessoria cadastrada (agregado
    global — especificação, "Cadastro de assessorias e apelidos": "nome
    que não está no cadastro vira aviso global")."""
    fontes = carregar("fontes")
    regras = carregar("regras")
    nomes_por_apelido = construir_nomes_por_apelido(db)

    wb_laudo = load_workbook(caminhos["laudo"], read_only=True, data_only=True)
    wb_iniciais = load_workbook(caminhos["iniciais"], read_only=True, data_only=True)
    wb_agendamento = load_workbook(caminhos["agendamento"], read_only=True, data_only=True)
    wb_extrajudicial = load_workbook(caminhos["extrajudicial"], read_only=True, data_only=True)
    wb_contrarias = load_workbook(caminhos["contrarias"], read_only=True, data_only=True)
    wb_procon = load_workbook(caminhos["procon"], read_only=True, data_only=True)

    linhas_laudo = aba_mensal.ler(wb_laudo, "laudo", fontes["laudos"], mes=mes, ano=ano)

    abas_do_ano = aba_mensal.encontrar_todas_abas_do_ano(wb_iniciais.sheetnames, ano=ano)
    linhas_iniciais_todas = []
    for mes_aba in abas_do_ano:
        linhas_iniciais_todas.extend(aba_mensal.ler(wb_iniciais, "iniciais", fontes["iniciais"], mes=mes_aba, ano=ano))

    linhas_judicial = aba_unica.ler(wb_agendamento, "agendamento", fontes["audiencias_judiciais"])
    linhas_pauta = aba_unica.ler(wb_agendamento, "agendamento", fontes["audiencias_contrarias"])
    linhas_extrajudicial = aba_unica.ler(wb_extrajudicial, "extrajudicial", fontes["extrajudiciais"])

    por_assessoria_contrarias = aba_por_assessoria.ler(wb_contrarias, "contrarias", fontes["contrarias"], nomes_por_apelido)
    por_assessoria_procon = aba_por_assessoria.ler(wb_procon, "procon", fontes["procons"], nomes_por_apelido)

    empresas_desconhecidas: dict[str, int] = {}
    for linhas, campo in (
        (linhas_laudo, "empresa"), (linhas_iniciais_todas, "assessoria"),
        (linhas_judicial, "empresa"), (linhas_pauta, "empresa"), (linhas_extrajudicial, "empresa"),
    ):
        for desconhecida in validacao.detectar_empresas_desconhecidas(linhas, campo, nomes_por_apelido):
            empresas_desconhecidas[desconhecida.nome] = empresas_desconhecidas.get(desconhecida.nome, 0) + desconhecida.quantidade

    nomes_assessorias = list(
        db.scalars(select(EmpresaCliente.nome).where(EmpresaCliente.ativo.is_(True)).order_by(EmpresaCliente.nome))
    )

    resultados = []
    for nome_assessoria in nomes_assessorias:
        linhas_laudo_a = filtrar_linhas(linhas_laudo, "empresa", nome_assessoria, nomes_por_apelido)
        linhas_iniciais_a = filtrar_linhas(linhas_iniciais_todas, "assessoria", nome_assessoria, nomes_por_apelido)
        linhas_judicial_a = filtrar_linhas(linhas_judicial, "empresa", nome_assessoria, nomes_por_apelido)
        linhas_pauta_a = filtrar_linhas(linhas_pauta, "empresa", nome_assessoria, nomes_por_apelido)
        linhas_extra_a = filtrar_linhas(linhas_extrajudicial, "empresa", nome_assessoria, nomes_por_apelido)
        linhas_contrarias_a = por_assessoria_contrarias.get(nome_assessoria, [])
        linhas_procon_a = por_assessoria_procon.get(nome_assessoria, [])

        r_laudos = laudos.calcular(linhas_laudo_a, data_corte, regras["laudos"])
        r_iniciais = iniciais.calcular(linhas_iniciais_a, mes=mes, ano=ano, data_corte=data_corte, regras=regras["iniciais"])
        r_extra = extrajudiciais.calcular(linhas_extra_a, mes=mes, ano=ano, data_corte=data_corte, regras=regras["extrajudiciais"])
        r_judicial = judiciais.calcular(linhas_judicial_a, ano=ano, data_corte=data_corte)
        r_pauta = judiciais.calcular(linhas_pauta_a, ano=ano, data_corte=data_corte)
        r_contrarias = contrarias.calcular(linhas_contrarias_a, mes=mes, ano=ano, data_corte=data_corte, regras=regras["contrarias"])
        r_procons = procons.calcular(linhas_procon_a, mes=mes, ano=ano, regras=regras["procons"])

        avisos: list[Aviso] = []
        for resultado_secao in (r_laudos, r_iniciais, r_extra, r_judicial, r_pauta, r_contrarias, r_procons):
            avisos.extend(resultado_secao.avisos)
        avisos.extend(validacao.conferir_cronologia_laudos(linhas_laudo_a))
        avisos.extend(validacao.conferir_cronologia_iniciais(linhas_iniciais_a))
        aviso_soma_uf = validacao.conferir_soma_uf(r_contrarias)
        if aviso_soma_uf is not None:
            avisos.append(aviso_soma_uf)
        avisos.extend(validacao.conferir_campos_essenciais(linhas_iniciais_a, ["data_recebimento", "numero_processo", "estados"]))
        avisos.extend(validacao.conferir_campos_essenciais(linhas_judicial_a, ["data", "processo"]))
        avisos.extend(validacao.conferir_campos_essenciais(linhas_pauta_a, ["data", "processo", "estado"]))

        dados = DadosRelatorio(
            assessoria=nome_assessoria,
            mes=mes,
            ano=ano,
            data_corte=data_corte,
            laudos=r_laudos,
            iniciais=r_iniciais,
            extrajudiciais=r_extra,
            audiencias_judiciais=r_judicial,
            audiencias_contrarias=r_pauta,
            contrarias=r_contrarias,
            procons=r_procons,
            manuais=manuais.CamposManuais(),
            avisos=avisos,
        )
        resultados.append(armazenamento.salvar_resultado(db, dados, usuario))

    lista_desconhecidas = sorted(
        (EmpresaDesconhecidaGlobal(nome, quantidade) for nome, quantidade in empresas_desconhecidas.items()),
        key=lambda e: -e.quantidade,
    )
    return resultados, lista_desconhecidas
